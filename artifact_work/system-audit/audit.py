"""Read-only product audit: isolated storage, real routes/parser/depot/templates."""
import copy, io, json, os, shutil, sys, unittest
from pathlib import Path
from unittest.mock import patch
from concurrent.futures import ThreadPoolExecutor
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT))
import depo, excel_araclari as ex
from tests.test_excel_sync import SyncTests
from tests.test_ui import UITests
from flask import template_rendered
from werkzeug.security import generate_password_hash
import openpyxl

PACK=ROOT/'outputs/pdgm-import-test-20260920'
OUT=ROOT/'outputs/audit-fixes-20260922'
OUT.mkdir(parents=True,exist_ok=True)
PASSWORD='Audit-Local-Only-2026!'
PASSWORD_HASH=generate_password_hash(PASSWORD)
OBS={}

class Audit(unittest.TestCase):
    def setUp(self):
        SyncTests.setUp(self)
        self.mod,self.client=UITests.load_app(self)
        self.mod.app.static_folder=str(ROOT/'static')
        self.users={'tester':{'rol':'admin','ad':'Test Yöneticisi','aktif':True,'sifre_hash':PASSWORD_HASH}}
        for user,typ in [('machine','makine'),('hand','elle_dizgi'),('eum','eum_dizgi')]:
            self.users[user]={'rol':'operator','ad':user,'operator_tipi':typ,'aktif':True,'sifre_hash':PASSWORD_HASH}
        self.users['viewer']={'rol':'gozlemci','ad':'Test Gözlemci','aktif':True,'sifre_hash':PASSWORD_HASH}
        self.save_users()
        self.upload_file('01_ILK_IMPORT.xlsx')

    def save_users(self):
        Path(self.mod.KULLANICI_DOSYASI).write_text(json.dumps(self.users),encoding='utf-8')
        self.mod._kullanici_mtime=None;self.mod._kullanici_onbellek=None

    def as_user(self,name):
        c=self.mod.app.test_client()
        with c.session_transaction() as s:s.update(kullanici=name,rol=self.users[name]['rol'],ad=name,operator_tipi=self.users[name].get('operator_tipi'),csrf_token='test-token')
        return c

    def api(self,path,data,client=None):
        if path=='/api/admin/duzenle' and 'surum' not in data:
            data={**data,'surum':depo.kart_getir(data['kart_id'])['surum']}
        return (client or self.client).post(path,json=data,headers={'X-CSRF-Token':'test-token'})

    def snap(self,source):
        target=self.root/'snapshot.xlsx';shutil.copy2(source,target);return str(target)

    def upload_file(self,name,via_route=False):
        p=PACK/name
        with patch.object(ex,'excel_deger_snapshot_olustur',self.snap):
            if via_route:
                with p.open('rb') as f:return self.client.post('/yonetim/yukle',data={'_csrf_token':'test-token','dosya':(f,name)})
            return ex.excelden_aktar(str(p),'tester')

    def parsed(self):
        with patch.object(ex,'excel_deger_snapshot_olustur',self.snap),patch.object(depo,'excel_import_uygula',side_effect=lambda **kw:kw):
            return ex.excelden_aktar(str(PACK/'01_ILK_IMPORT.xlsx'),'tester')

    def card(self,no,sheet='MAKINE'):
        return next(k for k in depo._kartlar if k.get('source_sheet')==sheet and k.get('source_row_id')=='NO:'+str(no))

    def progress(self,n=3):
        k=self.card(1012)
        self.assertEqual(self.api('/api/basla',{'kart_id':k['id'],'adet':10,'not':'Başlangıç notu'}).status_code,200)
        self.assertEqual(self.api('/api/bitir',{'kart_id':k['id'],'adet':n,'not':'Üretim notu'}).status_code,200)
        return k

    def test_01_real_upload_routes_update_api_templates_disk(self):
        ids={k['source_key']:k['id'] for k in depo._kartlar}
        self.assertEqual(self.upload_file('02_YENIDEN_IMPORT.xlsx',True).status_code,302)
        self.assertEqual(self.card(1005)['plan_teslim'],'2026-10-08')
        self.assertEqual(self.card(1007)['durum'],'DİZGİDE')
        self.assertEqual(self.card(1003)['source_active'],0)
        self.upload_file('03_TESLIM_VE_GERI_GELEN.xlsx',True)
        self.assertEqual(self.card(1007)['durum'],'TESLİM EDİLDİ')
        self.assertEqual(self.card(1003)['source_active'],1)
        self.assertTrue(all(k['id']==ids[k['source_key']] for k in depo._kartlar if k['source_key'] in ids))
        self.client.post('/yonetim/yeniden-oku',data={'_csrf_token':'test-token'})
        cards=self.client.get('/api/veriler').json['kartlar']
        self.assertEqual(next(k for k in cards if k['source_row_id']=='NO:1007')['durum'],'TESLİM EDİLDİ')

    def test_02_progress_notes_blank_status_restart(self):
        k=self.progress();kid=k['id']
        self.api('/api/not',{'kart_id':kid,'not':'Üçüncü not\nİkinci satır','isim':'Deneme Çığ'})
        note=k['aciklama'];self.upload_file('02_YENIDEN_IMPORT.xlsx');self.upload_file('03_TESLIM_VE_GERI_GELEN.xlsx')
        self.assertEqual(self.card(1012)['tamamlanan_adet'],3);self.assertEqual(self.card(1012)['aciklama'],note)
        depo.kur();self.assertEqual(self.card(1012)['aciklama'],note)
        self.assertIn('Üçüncü not',self.client.get('/operator').get_data(as_text=True))

    def test_03_excel_valid_status_quantity_matrix(self):
        expected={depo.PLANA_ALINDI:0,depo.HAZIR:0,depo.DIZGIDE:3,depo.TESLIM_EDILDI:10}
        original=self.parsed()
        for state,qty in expected.items():
            with self.subTest(state=state):
                self.upload_file('01_ILK_IMPORT.xlsx');k=self.progress();note=k['aciklama']
                p=copy.deepcopy(original);r=next(r for r in p['satirlar'] if r['source_row_id']=='NO:1012');r['ilk_durum']=state;r['plan']['excel_durum']=state
                depo.excel_import_uygula(**p)
                self.assertEqual(k['tamamlanan_adet'],qty);self.assertEqual(k['aciklama'],note);self.assertEqual(k['durum'],state)
        OBS['quantity_policy']={k:v for k,v in expected.items()}

    def test_04_full_operator_flow_all_three_types(self):
        for user,no,sheet in [('machine',1012,'MAKINE'),('hand',2007,'ELLE'),('eum',3003,'EUM')]:
            with self.subTest(user=user):
                c=self.as_user(user);k=self.card(no,sheet);kid=k['id'];total=k['toplam_adet']
                self.assertEqual(self.api('/api/basla',{'kart_id':kid,'adet':total},c).status_code,200)
                self.assertEqual(self.api('/api/teslim-et',{'kart_id':kid},c).status_code,409)
                for n in [0,-1,total+1]:self.assertEqual(self.api('/api/bitir',{'kart_id':kid,'adet':n},c).status_code,409)
                r=self.api('/api/bitir',{'kart_id':kid,'adet':total},c)
                self.assertTrue(r.json['uretim_bitti']);self.assertEqual(k['durum'],'DİZGİDE')
                self.assertEqual(self.api('/api/teslim-et',{'kart_id':kid,'not':'Teslim notu'},c).status_code,200)
                self.assertEqual(self.api('/api/teslim-et',{'kart_id':kid},c).status_code,409)
                self.assertTrue(k['teslim_zamani']);self.assertIn('Teslim notu',k['aciklama'])

    def test_05_roles_csrf_and_cross_type_writes(self):
        for user in ('machine','hand','eum','viewer'):
            self.assertEqual(self.as_user(user).get('/yonetim').status_code,403)
            self.assertEqual(self.api('/api/admin/kart-ekle',{},self.as_user(user)).status_code,403)
        self.assertEqual(self.as_user('viewer').get('/operator').status_code,403)
        self.assertEqual(self.client.post('/api/not',json={'kart_id':1,'not':'x'}).status_code,403)
        self.assertEqual(self.mod.app.test_client().get('/api/veriler').status_code,401)
        for path,data in [('/api/basla',{'kart_id':self.card(2007,'ELLE')['id'],'adet':1}),('/api/bitir',{'kart_id':self.card(2005,'ELLE')['id'],'adet':1}),('/api/teslim-et',{'kart_id':self.card(2005,'ELLE')['id']})]:
            self.assertEqual(self.api(path,data,self.as_user('machine')).status_code,409)

    def test_06_inactive_hidden_and_admin_restore(self):
        kid=self.card(1003)['id'];self.upload_file('02_YENIDEN_IMPORT.xlsx')
        self.assertEqual(self.api('/api/not',{'kart_id':kid,'not':'x'},self.as_user('machine')).status_code,409)
        self.assertEqual(self.api('/api/bitir',{'kart_id':kid,'adet':1},self.as_user('machine')).status_code,409)
        self.assertEqual(self.api('/api/not',{'kart_id':kid,'not':'Audit notu'}).status_code,200)
        self.upload_file('03_TESLIM_VE_GERI_GELEN.xlsx');self.assertIn('Audit notu',self.card(1003)['aciklama'])
        self.api('/api/admin/kart-sil',{'kart_id':kid});self.upload_file('03_TESLIM_VE_GERI_GELEN.xlsx')
        self.assertEqual(self.card(1003)['admin_gizli'],1)
        self.client.post('/yonetim/kart-geri-getir',data={'_csrf_token':'test-token','kart_id':kid})
        self.assertEqual(self.card(1003)['admin_gizli'],0)

    def test_07_material_confirmation(self):
        k=self.card(1013)
        self.assertEqual(self.api('/api/admin/duzenle',{'kart_id':k['id'],'durum':'PLANA ALINDI','toplam_adet':5,'tamamlanan_adet':0}).status_code,200)
        c=self.as_user('machine')
        r=self.api('/api/basla',{'kart_id':k['id'],'adet':5},c)
        self.assertEqual(r.status_code,409);self.assertTrue(r.json['malzeme_bekliyor_onay_gerekli'])
        self.assertEqual(self.api('/api/basla',{'kart_id':k['id'],'adet':5,'malzeme_onayi':True},c).status_code,200)
        self.assertEqual(k['malzeme_bekliyor'],0)

    def test_08_backup_restore_and_manual_card(self):
        k=self.progress();old=copy.deepcopy(k)
        backup=depo.anlik_yedek('audit')
        self.upload_file('05_TUM_KAYNAKLAR_BOS.xlsx')
        r=self.client.post('/yonetim/yedek-geri-yukle',data={'_csrf_token':'test-token','yedek':Path(backup).name})
        self.assertEqual(r.status_code,302);self.assertEqual(self.card(1012)['tamamlanan_adet'],old['tamamlanan_adet'])
        self.assertEqual(self.card(1012)['aciklama'],old['aciklama'])
        r=self.api('/api/admin/kart-ekle',{'talep_no':'MANUEL-TEST','stok_no':'MANUEL-STOK','toplam_adet':7,'not':'Manuel not'})
        self.assertEqual(r.status_code,201);kid=r.json['kart']['id'];self.upload_file('05_TUM_KAYNAKLAR_BOS.xlsx')
        self.assertEqual([k['id'] for k in depo.kartlari_getir()],[kid])

    def test_09_reports_persistence_and_escaping(self):
        kid=self.card(1012)['id'];payload='<img src=x onerror=alert(1)> & Türkçe'
        self.api('/api/not',{'kart_id':kid,'not':payload})
        for url in ['/operator','/yonetim']:
            html=self.client.get(url).get_data(as_text=True);self.assertNotIn(payload,html);self.assertIn('&lt;img',html)
        for url in ['/yonetim/rapor','/yonetim/kayit-dosyasi/kartlar','/yonetim/kayit-dosyasi/log','/yonetim/kayit-dosyasi/yuklemeler']:
            response=self.client.get(url);self.assertEqual(response.status_code,200)
            w=openpyxl.load_workbook(io.BytesIO(response.data),read_only=True,data_only=False)
            self.assertTrue(w.sheetnames);w.close();response.close()
        self.api('/api/not',{'kart_id':kid,'not':'=1+1'})
        w=openpyxl.load_workbook(depo.KARTLAR_DOSYA,read_only=True,data_only=False)
        self.assertFalse(any(c.data_type=='f' for ws in w for row in ws for c in row));w.close()

    def test_10_delivery_filters_and_invalid_ranges(self):
        for typ in ['MAKINE','ELLE','EUM']:
            r=self.client.get('/api/panel/teslimler?dizgi='+typ).json
            self.assertTrue(all(k['dizgi_kod']==typ for k in r['teslim_edilen']))
        for q in ['aralik=ozel&baslangic=bad&bitis=2026-01-01','aralik=ozel&baslangic=2026-12-31&bitis=2026-01-01']:
            self.assertEqual(self.client.get('/api/panel/teslimler?'+q).status_code,400)
        r=self.client.get('/api/panel/teslimler?aralik=ozel&baslangic=2026-09-01&bitis=2026-09-30').json
        self.assertEqual(len(r['teslim_edilen']),r['ozet']['kart'])
        self.assertNotIn('TEST-1008',[k['talep_no'] for k in r['teslim_edilen']])

    def test_11_monitor_and_badges(self):
        contexts=[]
        def capture(sender,template,context,**kw):contexts.append(context)
        with template_rendered.connected_to(capture,self.mod.app):r=self.client.get('/monitor')
        self.assertEqual({k['dizgi_kod'] for k in contexts[-1]['dizgide']+contexts[-1]['plana_alindi']},{'MAKINE'})
        self.assertIn('>PLANDA<',r.get_data(as_text=True));self.assertIn('PLANINDA (',self.client.get('/panel').get_data(as_text=True))
        with patch.object(depo,'bugun',return_value='2026-09-20'):
            for no,badge in [(1015,'SON GÜN'),(1016,'SON 1 GÜN'),(1019,'BUGÜN BAŞLAMALI')]:self.assertEqual(depo.kart_gorunumu(self.card(no))['rozet'],badge)

    def test_12_auth_login_logout_disable_and_redirect(self):
        c=self.mod.app.test_client()
        self.assertEqual(c.post('/giris',data={'kullanici':'machine','sifre':'invalid'}).status_code,401)
        r=c.post('/giris?devam=https://example.com',data={'kullanici':'machine','sifre':PASSWORD});self.assertEqual(r.status_code,302);self.assertNotIn('example.com',r.location)
        with c.session_transaction() as s:token=s['csrf_token']
        self.assertEqual(c.get('/operator').status_code,200)
        self.assertEqual(c.post('/cikis',data={'_csrf_token':token}).status_code,302)
        self.assertEqual(c.get('/operator').status_code,302)
        c=self.as_user('machine');self.users['machine']['aktif']=False;self.save_users()
        self.assertEqual(c.get('/api/veriler').status_code,401)

    def test_13_atomic_failure_import_and_operator(self):
        k=self.card(1012);before=copy.deepcopy(depo._kartlar);disk=Path(depo.KARTLAR_DOSYA).read_bytes()
        with patch.object(depo,'_coklu_yaz',side_effect=PermissionError('audit simulated lock')):
            self.assertEqual(self.api('/api/basla',{'kart_id':k['id'],'adet':10}).status_code,423)
        self.assertEqual(depo._kartlar,before);self.assertEqual(Path(depo.KARTLAR_DOSYA).read_bytes(),disk)

    def test_14_simultaneous_notes_preserved(self):
        kid=self.card(1012)['id']
        with ThreadPoolExecutor(max_workers=4) as pool:list(pool.map(lambda n:depo.kart_not_guncelle(kid,'Concurrent-'+str(n),'tester','admin'),range(8)))
        note=self.card(1012)['aciklama']
        for n in range(8):self.assertEqual(note.count('Concurrent-'+str(n)),1)

    def test_15_empty_wrong_extension_and_corrupt_uploads(self):
        before=copy.deepcopy(depo._kartlar)
        for data in [{'_csrf_token':'test-token'},{'_csrf_token':'test-token','dosya':(io.BytesIO(b'abc'),'wrong.txt')},{'_csrf_token':'test-token','dosya':(io.BytesIO(b'abc'),'bad.xlsx')}]:
            with patch.object(ex,'excel_deger_snapshot_olustur',self.snap):
                r=self.client.post('/yonetim/yukle',data=data)
            self.assertEqual(r.status_code,302);self.assertEqual(depo._kartlar,before)

    def test_16_identity_different_sheets_duplicate_business_keys(self):
        self.assertEqual(self.card(1001)['toplam_adet'],5);self.assertEqual(self.card(1001,'ELLE')['toplam_adet'],2)
        self.assertEqual(self.card(1001)['durum'],'DİZGİDE');self.assertEqual(self.card(1001,'ELLE')['durum'],'TESLİM EDİLDİ')
        self.assertEqual(len({self.card(n)['id'] for n in [1002,1003,1004]}),3)

    def test_17_admin_same_status_note_does_not_fabricate_date(self):
        k=self.card(1008)
        r=self.api('/api/admin/duzenle',{'kart_id':k['id'],'durum':k['durum'],'toplam_adet':5,'tamamlanan_adet':5,'not':'Yalnız not'})
        self.assertEqual(r.status_code,200)
        for field in ['gerceklesen_teslim','teslim_zamani','baslama_zamani','bitis_zamani']:self.assertIsNone(k[field])

    # Desired-behavior checks below intentionally remain red if a product defect exists.
    def test_bug_01_all_time_summary_matches_delivered_rows(self):
        r=self.client.get('/api/panel/teslimler?aralik=tumu').json
        OBS['all_time_summary']={'rows':len(r['teslim_edilen']),'summary_cards':r['ozet']['kart']}
        self.assertEqual(r['ozet']['kart'],len(r['teslim_edilen']),'Tümü summary excludes undated delivered cards')

    def test_bug_02_json_fractional_completion_rejected(self):
        k=self.progress();r=self.api('/api/bitir',{'kart_id':k['id'],'adet':1.9})
        OBS['fractional_quantity']={'status':r.status_code,'completed':k['tamamlanan_adet']}
        self.assertGreaterEqual(r.status_code,400,'JSON 1.9 silently truncated to 1')

    def test_bug_03_admin_stale_form_does_not_erase_new_note(self):
        k=self.progress();stale=copy.deepcopy(k)
        self.api('/api/not',{'kart_id':k['id'],'not':'YENİ OPERATÖR NOTU'})
        r=self.api('/api/admin/duzenle',{'kart_id':k['id'],'surum':depo.kart_surumu(stale),'durum':stale['durum'],'toplam_adet':stale['toplam_adet'],'tamamlanan_adet':stale['tamamlanan_adet'],'not':stale['aciklama']})
        OBS['stale_admin_note']={'status':r.status_code,'new_note_preserved':'YENİ OPERATÖR NOTU' in k['aciklama']}
        self.assertTrue(r.status_code==409 or 'YENİ OPERATÖR NOTU' in k['aciklama'],'Stale admin form overwrote newer note')

    def test_bug_04_live_operator_type_change_takes_effect(self):
        c=self.as_user('machine');self.users['machine']['operator_tipi']='elle_dizgi';self.save_users()
        r=self.api('/api/basla',{'kart_id':self.card(1012)['id'],'adet':10},c)
        OBS['stale_operator_type']={'status':r.status_code,'file_type':'elle_dizgi'}
        self.assertEqual(r.status_code,409,'Existing session retains old operator type')

    def test_bug_05_panel_type_filter_has_all_planned_cards(self):
        contexts=[]
        def capture(sender,template,context,**kw):contexts.append(context)
        with template_rendered.connected_to(capture,self.mod.app):self.client.get('/panel')
        actual=[k['id'] for k in depo.kartlari_getir() if k['durum']=='PLANA ALINDI']
        sent=[k['id'] for k in contexts[-1]['plana_alindi']]
        OBS['planned_truncation']={'all':len(actual),'sent_to_panel':len(sent)}
        self.assertEqual(set(actual),set(sent),'Panel slices planned cards to first 12 before browser filters')

    def test_bug_06_stale_admin_form_does_not_reduce_completed(self):
        k=self.progress();stale=copy.deepcopy(k);self.api('/api/bitir',{'kart_id':k['id'],'adet':2})
        r=self.api('/api/admin/duzenle',{'kart_id':k['id'],'surum':depo.kart_surumu(stale),'durum':stale['durum'],'toplam_adet':10,'tamamlanan_adet':stale['tamamlanan_adet'],'not':stale['aciklama']})
        OBS['stale_admin_quantity']={'status':r.status_code,'completed_after':k['tamamlanan_adet'],'expected':5}
        self.assertTrue(r.status_code==409 or k['tamamlanan_adet']==5,'Stale form silently reduced 5 completed to 3')

    def test_bug_07_oversized_note_not_silently_truncated(self):
        kid=self.card(1012)['id'];note='A'*33000+'END-MARKER'
        r=self.api('/api/not',{'kart_id':kid,'not':note})
        if r.status_code>=400:return
        before=self.card(1012)['aciklama'];depo.kur();after=self.card(1012)['aciklama']
        OBS['long_note']={'status':r.status_code,'memory_length':len(before),'disk_length':len(after)}
        self.assertEqual(before,after,'Excel maximum cell text length silently loses note tail')

if __name__=='__main__':
    if '--serve' in sys.argv:
        case=Audit();case.setUp()
        case.mod.app.config['TESTING']=False
        print('AUDIT_SERVER http://127.0.0.1:5017 LOGIN tester / '+PASSWORD,flush=True)
        try:case.mod.app.run(host='127.0.0.1',port=5017,threaded=True,use_reloader=False)
        finally:case.doCleanups()
    else:
        suite=unittest.defaultTestLoader.loadTestsFromTestCase(Audit)
        with (OUT/'automated-tests.txt').open('w',encoding='utf-8') as log:
            result=unittest.TextTestRunner(stream=log,verbosity=2).run(suite)
        report={'tests':result.testsRun,'failures':[{'test':str(t),'traceback':e[-4000:]} for t,e in result.failures],'errors':[{'test':str(t),'traceback':e[-4000:]} for t,e in result.errors],'observations':OBS}
        (OUT/'results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
        print(json.dumps(report,ensure_ascii=False,indent=2))
        sys.exit(0 if result.wasSuccessful() else 1)
