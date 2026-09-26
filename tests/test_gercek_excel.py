"""Gerçek PDGM Excel'i (ve değiştirilmiş kopyaları) ile uçtan uca import senaryoları.

Üretim data klasörüne dokunulmaz: depo yolları her testte TemporaryDirectory'ye
yönlenir, kaynak Excel yalnız okunur, değişiklikler geçici kopyalarda yapılır.
Varsayılan koşu COM yerine dosyanın kendisini snapshot sayar (dosyada formül
yok; COM ile birebir aynı sonuç ayrıca doğrulanır). PDGM_TEST_EXCEL_COM=1 ile
aynı senaryoların bir bölümü gerçek Microsoft Excel üzerinden de koşar.
"""
import copy
import os
import shutil
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest.mock import patch

import openpyxl
from openpyxl.utils.datetime import from_excel

import depo
import excel_araclari as ex
from tests import test_excel_sync as fixtures, test_ui as ui
from tests.test_excel_sync import ROOT, row

GERCEK_EXCEL = ROOT / 'PDGM_Kart_dizgi_Talepleri_Üretim_Takvimi.xlsx'

# Parser'dan bağımsız referans: bu dosyanın sabit sütun harfleri.
# (sayfa, kod, başlık satırı, NO, Talep, Stok, Adet, Dizgi Başlama, Plan Teslim, Gerçekleşen, DURUM)
DUZEN = [
    ('MAKİNE', 'MAKINE', 4, 'C', 'D', 'I', 'J', 'U', 'W', 'X', 'Y'),
    ('ELDE DİZGİ', 'ELLE', 2, 'B', 'C', 'H', 'I', 'S', 'T', 'U', 'V'),
    ('EÜM', 'EUM', 2, 'B', 'C', 'H', 'I', None, 'U', 'V', 'W'),
]
BEKLENEN_DURUM = {
    'TESLİM EDİLDİ': 'TESLİM EDİLDİ', 'DİZGİDE': 'DİZGİDE', 'PLANA ALINDI': 'PLANA ALINDI',
    'DİZGİDE VE MALZEME BEKLENİYOR': 'DİZGİDE', 'DİZGİ İÇİN BEKLİYOR': 'PLANA ALINDI',
    'ÜRETİM DEVAM EDİYOR': 'DİZGİDE', 'MALZEME TEDARİK': None, 'PDGM ÖNERİ': None, None: None,
}
# EÜM'de başlangıç = "NN. hafta (G.AA haftası)" metninin Pazartesi'si (elle hesaplandı).
EUM_BASLANGIC = {1: '2025-06-02', 2: '2025-06-09', 3: '2025-06-16', 4: '2025-06-23', 5: '2025-06-23',
                 6: '2025-06-23', 7: '2026-07-27', 8: '2026-09-07', 9: '2026-09-07'}
MAKINE_KALINTI = '211–225, 229–253, 255–259, 262, 265–268, 270–271'
EUM_KALINTI = '10, 205, 208, 211'


def _bos(v):
    return v is None or (isinstance(v, str) and not v.strip())


def _iso(v):
    if _bos(v):
        return None
    if isinstance(v, (int, float)):
        v = from_excel(v)
    return v.strftime('%Y-%m-%d')


def referans_kayitlar(path):
    """Talep NO ve Kart Stok No dolu satırlar: {(kod, NO): beklenen alanlar}."""
    wb = openpyxl.load_workbook(path, data_only=True)
    sonuc = {}
    for sayfa, kod, baslik, c_no, c_talep, c_stok, c_adet, c_bas, c_tes, c_ger, c_durum in DUZEN:
        ws = wb[sayfa]
        for r in range(baslik + 1, ws.max_row + 1):
            talep, stok = ws[f'{c_talep}{r}'].value, ws[f'{c_stok}{r}'].value
            if _bos(talep) or _bos(stok):
                continue
            no = int(ws[f'{c_no}{r}'].value)
            durum_ham = ws[f'{c_durum}{r}'].value
            durum_ham = None if _bos(durum_ham) else durum_ham.strip()
            sonuc[(kod, no)] = {
                'talep_no': str(talep).strip(), 'stok_no': str(stok).strip(),
                'toplam_adet': int(str(ws[f'{c_adet}{r}'].value).split()[0]),
                'plan_baslama': EUM_BASLANGIC[no] if kod == 'EUM' else _iso(ws[f'{c_bas}{r}'].value),
                'plan_teslim': _iso(ws[f'{c_tes}{r}'].value),
                'gerceklesen_teslim': _iso(ws[f'{c_ger}{r}'].value),
                'excel_durum': durum_ham, 'durum': BEKLENEN_DURUM[durum_ham], 'satir': r,
            }
    wb.close()
    return sonuc


@unittest.skipUnless(GERCEK_EXCEL.exists(), 'Gerçek PDGM Excel dosyası proje kökünde yok')
class GercekExcelTests(unittest.TestCase):
    setUp = fixtures.SyncTests.setUp
    load_app = ui.UITests.load_app

    # --- yardımcılar -----------------------------------------------------
    def kopya(self, degistir=None, ad='kopya.xlsx'):
        wb = openpyxl.load_workbook(GERCEK_EXCEL)
        if degistir:
            degistir(wb)
        path = self.root / ad
        wb.save(path)
        wb.close()
        return path

    def snapshot_yamasi(self):
        def snapshot(kaynak):
            hedef = self.root / 'snapshot.xlsx'
            shutil.copy2(kaynak, hedef)
            return str(hedef)
        return patch.object(ex, 'excel_deger_snapshot_olustur', snapshot)

    def yukle(self, path=GERCEK_EXCEL, **kw):
        with self.snapshot_yamasi():
            return ex.excelden_aktar(str(path), 'tester', **kw)

    def kartlar(self):
        return {(k['source_sheet'], int(k['source_row_id'][3:])): k for k in depo._kartlar
                if k.get('kaynak') == 'EXCEL'}

    def disk(self):
        return {p: Path(p).read_bytes() for p in (depo.KARTLAR_DOSYA, depo.LOG_DOSYA, depo.YUKLEME_DOSYA)}

    # --- 1: orijinal dosya ------------------------------------------------
    def test_orijinal_excel_eksiksiz_ve_dogru_aktarilir(self):
        beklenen = referans_kayitlar(GERCEK_EXCEL)
        self.assertEqual(len(beklenen), 44)
        sonuc = self.yukle()
        rapor = sonuc['kaynak_raporu']
        sayfa = {s['kod']: s for s in rapor['sayfalar']}
        self.assertEqual({k: s['kayit'] for k, s in sayfa.items()}, {'MAKINE': 15, 'ELLE': 20, 'EUM': 9})
        self.assertEqual(sayfa['MAKINE']['atlanan_araliklar'], MAKINE_KALINTI)
        self.assertEqual(sayfa['EUM']['atlanan_araliklar'], EUM_KALINTI)
        self.assertEqual(sayfa['ELLE']['atlanan'], [])
        self.assertTrue(all(a['gizli'] for s in sayfa.values() for a in s['atlanan']))
        self.assertEqual((sonuc['yeni'], sonuc['guncellenen'], sonuc['pasife_alinan']), (44, 0, 0))
        # Önizlemede uyarılar türüne göre gruplanır (9 uyarı -> 3 grup).
        self.assertEqual({(g['tur'], g['baslik']): (g['adet'], g['yerler']) for g in rapor['uyari_gruplari']}, {
            ('durum_bos', 'DURUM boş'): (4, [{'sayfa': 'MAKİNE', 'satirlar': '273–275, 277'}]),
            ('durum_durumsuz', 'DURUM "MALZEME TEDARİK": iş akışı durumu atanmaz'):
                (3, [{'sayfa': 'ELDE DİZGİ', 'satirlar': '67, 69, 72'}]),
            ('durum_durumsuz', 'DURUM "PDGM ÖNERİ": iş akışı durumu atanmaz'):
                (2, [{'sayfa': 'EÜM', 'satirlar': '244–245'}]),
        })
        self.assertEqual([(g['satirlar'], g['adet'], g['gizli']) for g in sayfa['EUM']['atlanan_gruplari']],
                         [('10', 1, 1), ('205, 208', 2, 2), ('211', 1, 1)])

        kartlar = self.kartlar()
        self.assertEqual(set(kartlar), set(beklenen))
        for anahtar, b in beklenen.items():
            k = kartlar[anahtar]
            with self.subTest(kart=anahtar):
                for alan in ('talep_no', 'stok_no', 'toplam_adet', 'plan_baslama', 'plan_teslim',
                             'gerceklesen_teslim', 'excel_durum', 'durum'):
                    self.assertEqual(k[alan], b[alan], alan)
                # Üretim türü yapısal olarak kaynak sayfadan gelir (ters görünme yok).
                self.assertEqual(depo.dizgi_kodu(k['dizgi_tipi']), anahtar[0])
                self.assertEqual(k['source_active'], 1)
                self.assertEqual(k['tamamlanan_adet'], k['toplam_adet'] if k['durum'] == 'TESLİM EDİLDİ' else 0)
        # Malzeme bayrağı: BEKLENİYOR yazımı ve MALZEME TEDARİK.
        self.assertEqual(sorted(n for (kod, n), k in kartlar.items() if k['malzeme_bekliyor']),
                         [13, 17, 18, 20])
        # Diskten yeniden yüklenen durum aynı.
        once = copy.deepcopy(depo._kartlar)
        depo._kartlar = []
        depo.kur()
        self.assertEqual(depo._kartlar, once)

    def test_ayni_talep_farkli_sayfada_bagimsiz(self):
        self.yukle()
        kartlar = self.kartlar()
        makine, eum = kartlar[('MAKINE', 11)], kartlar[('EUM', 8)]
        self.assertEqual(makine['talep_no'], eum['talep_no'])  # 1826751 iki sayfada da var
        self.assertNotEqual(makine['id'], eum['id'])

        def degistir(wb):
            wb['EÜM']['W244'] = 'ÜRETİM PLANA ALINDI'
        self.yukle(self.kopya(degistir))
        kartlar = self.kartlar()
        self.assertEqual(kartlar[('EUM', 8)]['durum'], 'PLANA ALINDI')
        self.assertEqual(kartlar[('MAKINE', 11)], makine)

    # --- 2-5: boş / gizli satır ve sütunlar -------------------------------
    def test_bos_ve_gizli_bos_satirlar_atlanir(self):
        def degistir(wb):
            ws = wb['MAKİNE']
            ws.insert_rows(226, 3)          # kayıtların arasına 3 boş satır
            for sutun in 'DEGIJY':
                ws[f'{sutun}227'] = '   '   # yalnız boşluk (formülün "" sonucu gibi)
            ws.row_dimensions[228].hidden = True
        sonuc = self.yukle(self.kopya(degistir))
        makine = next(s for s in sonuc['kaynak_raporu']['sayfalar'] if s['kod'] == 'MAKINE')
        self.assertEqual(makine['kayit'], 15)
        self.assertEqual(sonuc['yeni'], 44)
        self.assertEqual(makine['bos_satir'], 219 + 3)

    def test_gizli_dolu_satir_normal_islenir(self):
        def degistir(wb):
            ws = wb['MAKİNE']
            ws.row_dimensions[272].hidden = True          # mevcut kayıt (NO 11) gizlendi
            for sutun, deger in zip('CDEGIJUWY', [16, 1900001, 'Gizli Sahip', 'Ahmet Düzgün', 'AD-GIZLI-0001',
                                                  '7 ADET', datetime(2026, 9, 28), datetime(2026, 10, 2),
                                                  'PLANA ALINDI']):
                ws[f'{sutun}278'] = deger
            ws.row_dimensions[278].hidden = True          # yeni kayıt gizli satırda
        sonuc = self.yukle(self.kopya(degistir))
        makine = next(s for s in sonuc['kaynak_raporu']['sayfalar'] if s['kod'] == 'MAKINE')
        self.assertEqual((makine['kayit'], makine['gizli_kayit']), (16, 2))
        k = self.kartlar()
        self.assertEqual(k[('MAKINE', 11)]['durum'], 'PLANA ALINDI')
        self.assertEqual((k[('MAKINE', 16)]['stok_no'], k[('MAKINE', 16)]['toplam_adet'],
                          k[('MAKINE', 16)]['plan_baslama']), ('AD-GIZLI-0001', 7, '2026-09-28'))

    def test_gizli_ve_bos_sutunlar_baslik_eslesmesini_bozmaz(self):
        self.yukle()
        alanlar = ('durum', 'gerceklesen_teslim', 'plan_teslim', 'plan_baslama', 'pcb', 'stok_no',
                   'toplam_adet', 'dizgi_sorumlusu')
        once = {a: tuple(k[x] for x in alanlar) for a, k in self.kartlar().items()}
        gizlenecek = {'GERCEKLESEN TESLIM T', 'DURUM', 'PCB', 'MALZEME/PCB', 'KART STOK NO',
                      'KART URETIM ADET', 'DIZGI BASLAMA T'}

        def degistir(wb):
            # Başlıksız boş sütun: sonraki bütün sütunlar bir sağa kayar (sabit indeks yok).
            wb['MAKİNE'].insert_cols(5)
            for ad, baslik in (('MAKİNE', 4), ('ELDE DİZGİ', 2), ('EÜM', 2)):
                for hucre in wb[ad][baslik]:
                    if ex._sadelestir(hucre.value) in gizlenecek:
                        wb[ad].column_dimensions[hucre.column_letter].hidden = True
        sonuc = self.yukle(self.kopya(degistir))
        self.assertEqual((sonuc['yeni'], sonuc['guncellenen'], sonuc['pasife_alinan']), (0, 0, 0))
        sonra = {a: tuple(k[x] for x in alanlar) for a, k in self.kartlar().items()}
        self.assertEqual(sonra, once)

    def test_yinelenen_teslim_sutunu_gizlilikten_bagimsiz_uretim_blogundan_okunur(self):
        """Sipariş bloğundaki (O) ve üretim bloğundaki (W) "Planlanan Teslim T." aynı başlıkta.

        Eskiden görünür olan okunuyordu: W gizlenip O gösterilirse tarihler sessizce yanlış
        sütundan gelirdi. Artık DURUM'a yakın olan üretim sütunu okunur.
        """
        self.yukle(); once = copy.deepcopy(depo._kartlar)
        for aciklama, gizlilik in [("O görünür, W gizli", {'O': False, 'W': True}),
                                   ("ikisi de görünür", {'O': False, 'W': False}),
                                   ("ikisi de gizli", {'O': True, 'W': True})]:
            with self.subTest(aciklama):
                def degistir(wb):
                    for harf, gizli in gizlilik.items():
                        wb['MAKİNE'].column_dimensions[harf].hidden = gizli
                sonuc = self.yukle(self.kopya(degistir))
                self.assertEqual((sonuc['yeni'], sonuc['guncellenen'], sonuc['degismeyen']), (0, 0, 44))
                self.assertEqual(depo._kartlar, once)
                self.assertEqual(self.kartlar()[('MAKINE', 11)]['plan_teslim'], '2026-09-02')  # W, O değil

    def test_yinelenen_kimlik_sutunu_hala_belirsizse_hata(self):
        def degistir(wb):   # iki görünür NO sütunu: kimlik tahmin edilmez
            wb['MAKİNE']['V4'] = 'NO'
        with self.assertRaisesRegex(ex.ExcelAktarimHatasi, r"'NO' için birden fazla sütun bulundu \(C, V\)"):
            self.yukle(self.kopya(lambda wb: (degistir(wb), setattr(wb['MAKİNE'].column_dimensions['V'], 'hidden', False))))
        self.assertEqual(depo._kartlar, [])

    # --- 6, 13: eksik zorunlu alan / hatalı dosya ---------------------------
    def test_eksik_zorunlu_alan_anlasilir_hata_ve_veri_korunur(self):
        self.yukle()
        once, disk = copy.deepcopy(depo._kartlar), self.disk()
        yedekler = set(os.listdir(depo.YEDEK_KLASORU))

        def degistir(wb):
            wb['MAKİNE']['D272'] = None       # NO 11: Talep NO silindi, adet/DURUM/tarih duruyor
            wb['EÜM']['H9'] = '  '            # NO 7: Kart Stok No boşluk
            wb['ELDE DİZGİ']['I6'] = 'on adet'  # NO 1: adet okunamaz
        with self.assertRaises(ex.ExcelAktarimHatasi) as hata:
            self.yukle(self.kopya(degistir))
        mesaj = str(hata.exception)
        self.assertIn('3 sorun bulundu', mesaj)
        self.assertIn("'MAKİNE' satır 272: Talep NO boş, ancak satırda kayıt verisi var", mesaj)
        self.assertIn("'EÜM' satır 9: Kart Stok No boş", mesaj)
        self.assertIn("'ELDE DİZGİ' sayfası, satır 6: Üretim adedi", mesaj)
        self.assertEqual(depo._kartlar, once)
        self.assertEqual(self.disk(), disk)
        self.assertEqual(set(os.listdir(depo.YEDEK_KLASORU)), yedekler)  # depo'ya hiç ulaşılmadı

    def test_is_kurali_hatasinda_bellek_ve_disk_geri_alinir(self):
        self.yukle()
        kart = self.kartlar()[('MAKINE', 9)]            # DİZGİDE, 10 adet
        depo.kart_bitir(kart['id'], 4, 'op', 'operator', 'makine')
        once, disk = copy.deepcopy(depo._kartlar), self.disk()

        def degistir(wb):
            wb['MAKİNE']['J264'] = '3 ADET'            # tamamlanan (4) altına indi
            wb['EÜM']['W9'] = 'PLANA ALINDI'          # geçerli başka bir değişiklik de uygulanmamalı
        with self.assertRaises(depo.IsKuralHatasi):
            self.yukle(self.kopya(degistir))
        self.assertEqual(depo._kartlar, once)
        self.assertEqual(self.disk(), disk)

    # --- 7-8: EÜM ------------------------------------------------------------
    def test_eum_yeni_kayitlar_ve_plan_durumlari(self):
        self.yukle()

        def degistir(wb):
            ws = wb['EÜM']
            satirlar = [
                (246, 10, 'PLANA ALINDI', '39. hafta (21.09 haftası)', datetime(2026, 9, 25)),
                (247, 11, 'ÜRETİM PLANA ALINDI', '40. hafta (28.09 haftası)', None),
                (248, 12, 'üretim  plana alındı ', '41. Hafta (5.10 haftası)', datetime(2026, 10, 9)),
                (249, 13, 'PLANA ALİNDI', '42. hafta', datetime(2026, 10, 16)),
            ]
            for r, no, durum, hafta, teslim in satirlar:
                for s, v in zip('BCDHITUW', [no, 1900000 + no, 'EÜM Sahip', f'EU-0000-{no:04d}', '4 ADET',
                                           hafta, teslim, durum]):
                    ws[f'{s}{r}'] = v
            ws['W9'] = 'ÜRETİM PLANA ALINDI'           # DİZGİDE -> planlandı
            ws['W3'] = 'PLANA ALINDI'                  # TESLİM EDİLDİ -> planlandı (tarih dolu)
        sonuc = self.yukle(self.kopya(degistir))
        self.assertEqual(sonuc['yeni'], 4)
        k = self.kartlar()
        for no, baslangic in [(10, '2026-09-21'), (11, '2026-09-28'), (12, '2026-10-05'), (13, '2026-10-12')]:
            with self.subTest(no=no):
                kart = k[('EUM', no)]
                self.assertEqual((kart['dizgi_tipi'], kart['durum'], kart['tamamlanan_adet'],
                                  kart['baslangic_adet'], kart['plan_baslama']),
                                 ("EÜM'DE DİZGİ", 'PLANA ALINDI', 0, 0, baslangic))
        self.assertEqual((k[('EUM', 7)]['durum'], k[('EUM', 7)]['tamamlanan_adet']), ('PLANA ALINDI', 0))
        self.assertEqual((k[('EUM', 1)]['durum'], k[('EUM', 1)]['tamamlanan_adet']), ('PLANA ALINDI', 0))
        # Planlanmış EÜM kartı teslim/üretimde görünmez; EÜM filtresi karışmaz.
        mod, client = self.load_app()
        teslim = client.get('/api/panel/teslimler?dizgi=EUM').get_json()['teslim_edilen']
        self.assertEqual(len(teslim), 5)
        self.assertNotIn('1465221', {t['talep_no'] for t in teslim})
        pano = client.get('/api/veriler').get_json()
        eum_plan = [x for x in pano['kartlar'] if x['dizgi_kod'] == 'EUM' and x['durum'] == 'PLANA ALINDI']
        self.assertEqual(len(eum_plan), 6)
        self.assertEqual(pano['sayac']['plana_alindi_eum'], 6)
        self.assertEqual(pano['sayac']['dizgide_eum'], 0)
        self.assertTrue(all(x['rozet'] != 'TESLİM EDİLDİ' and 'TESLİM' not in x['rozet'] for x in eum_plan))

    # --- 9-12: güncelleme, ekleme, silme, tekrar yükleme --------------------
    def test_tarih_ve_durum_degisikligi_mevcut_karti_gunceller(self):
        mod, client = self.load_app()
        self.yukle()
        idler = {a: k['id'] for a, k in self.kartlar().items()}

        def degistir(wb):
            ws = wb['MAKİNE']
            ws['W272'] = datetime(2026, 9, 10)          # NO 11 plan teslim
            ws['Y272'] = 'DİZGİDE'                     # NO 11 PLANA -> DİZGİDE
            ws['X264'] = datetime(2026, 9, 1)           # NO 9 teslim edildi
            ws['Y264'] = 'TESLİM EDİLDİ'
            wb['ELDE DİZGİ']['S70'] = datetime(2026, 8, 25)  # ELDE NO 19 başlangıç
        sonuc = self.yukle(self.kopya(degistir))
        self.assertEqual((sonuc['yeni'], sonuc['guncellenen'], sonuc['degismeyen'], sonuc['pasife_alinan']),
                         (0, 3, 41, 0))
        k = self.kartlar()
        self.assertEqual({a: x['id'] for a, x in k.items()}, idler)
        self.assertEqual((k[('MAKINE', 11)]['plan_teslim'], k[('MAKINE', 11)]['durum']), ('2026-09-10', 'DİZGİDE'))
        self.assertEqual((k[('MAKINE', 9)]['durum'], k[('MAKINE', 9)]['gerceklesen_teslim'],
                          k[('MAKINE', 9)]['tamamlanan_adet']), ('TESLİM EDİLDİ', '2026-09-01', 10))
        self.assertEqual(k[('ELLE', 19)]['plan_baslama'], '2026-08-25')
        # API ve diskten yeniden yükleme eski veriyi göstermez.
        api = {(x['dizgi_kod'], x['sira']): x for x in client.get('/api/veriler').get_json()['kartlar']}
        self.assertEqual((api[('MAKINE', 11)]['plan_teslim'], api[('MAKINE', 11)]['durum']),
                         ('2026-09-10', 'DİZGİDE'))
        self.assertEqual(api[('MAKINE', 9)]['durum'], 'TESLİM EDİLDİ')
        depo._kartlar = []
        depo.kur()
        self.assertEqual(self.kartlar()[('MAKINE', 11)]['plan_teslim'], '2026-09-10')

    def test_yeni_kayit_eklenir_digerleri_degismez(self):
        self.yukle(); once = copy.deepcopy(self.kartlar())

        def degistir(wb):
            ws = wb['MAKİNE']
            for s, v in zip('CDEGIJTUWY', [16, 1850000, 'Yeni Sahip', 'Ahmet Düzgün', 'AD-YENI-0016', '12 ADET',
                                           '40. hafta (28.09 haftası)', datetime(2026, 9, 28),
                                           datetime(2026, 10, 1), 'PLANA ALINDI']):
                ws[f'{s}279'] = v
        sonuc = self.yukle(self.kopya(degistir))
        self.assertEqual((sonuc['yeni'], sonuc['guncellenen'], sonuc['pasife_alinan']), (1, 0, 0))
        sonra = self.kartlar()
        self.assertEqual({a: sonra[a] for a in once}, once)
        self.assertEqual((sonra[('MAKINE', 16)]['durum'], sonra[('MAKINE', 16)]['toplam_adet']), ('PLANA ALINDI', 12))

    def test_silinen_kayit_pasiflesir_geri_gelince_ayni_kart(self):
        self.yukle()
        kart = self.kartlar()[('MAKINE', 8)]

        def sil(wb):
            ws = wb['MAKİNE']
            for hucre in ws[263]:
                hucre.value = None
        onizleme = self.yukle(self.kopya(sil), onizleme=True)
        self.assertEqual(onizleme['pasife_alinan'], 1)
        self.assertEqual(self.kartlar()[('MAKINE', 8)]['source_active'], 1)   # önizleme yazmaz
        sonuc = self.yukle(self.kopya(sil))
        self.assertEqual((sonuc['pasife_alinan'], sonuc['yeni']), (1, 0))
        pasif = self.kartlar()[('MAKINE', 8)]
        self.assertEqual((pasif['id'], pasif['source_active']), (kart['id'], 0))
        self.assertNotIn(kart['id'], {k['id'] for k in depo.kartlari_getir()})
        self.assertEqual(len(depo.kartlari_getir()), 34)
        # Satır geri gelirse aynı kart geçmişiyle yeniden aktifleşir; çoğalma yok.
        sonuc = self.yukle()
        self.assertEqual((sonuc['yeni'], sonuc['guncellenen']), (0, 1))
        self.assertEqual(self.kartlar()[('MAKINE', 8)]['source_active'], 1)
        self.assertEqual(len(depo._kartlar), 44)

    def test_kalinti_birakilarak_silinen_kayit_raporlanir(self):
        self.yukle()

        def degistir(wb):  # NO 11: talep/adet/tarih/durum temizlendi; NO, stok, sorumlu, hafta kaldı
            for s in 'DFHJUWXY':
                wb['MAKİNE'][f'{s}272'] = None
        sonuc = self.yukle(self.kopya(degistir), onizleme=True)
        self.assertEqual(sonuc['pasife_alinan'], 1)
        makine = next(s for s in sonuc['kaynak_raporu']['sayfalar'] if s['kod'] == 'MAKINE')
        atlanan = {a['satir']: a for a in makine['atlanan']}
        self.assertIn(272, atlanan)
        self.assertEqual(atlanan[272]['dolu'], ['C:NO', 'E:Talep Sahibi', 'G:PDGM Dizgi Sorumlusu',
                                               'I:Kart Stok No', 'K:Birim Seviyesi Kullanım',
                                               'L:STROM İhtiyacı', 'M:Üretim Yeri', 'N:Sipariş Tarihi',
                                               'O:Planlanan Teslim T.', 'P:Gerçekleşen Teslim Tarihi',
                                               'Q:Üretim Yeri', 'R:Elek Durumu', 'S:Dizgi Talep Açılma Tarihi',
                                               'T:Planlanan Başlangıç T.'])

    def test_ayni_dosya_tekrar_yukleme_degisiklik_uretmez(self):
        self.yukle()
        once = copy.deepcopy(depo._kartlar)
        for yeniden_baslat in (False, True):
            with self.subTest(yeniden_baslat=yeniden_baslat):
                if yeniden_baslat:
                    depo._kartlar = []
                    depo.kur()
                log_sayisi = len(depo._loglar)
                onizleme = self.yukle(onizleme=True)
                self.assertEqual(onizleme['degisiklikler'], [])
                sonuc = self.yukle()
                self.assertEqual((sonuc['yeni'], sonuc['guncellenen'], sonuc['degismeyen'],
                                  sonuc['pasife_alinan']), (0, 0, 44, 0))
                self.assertEqual(depo._kartlar, once)  # guncelleme damgası ve sürüm dahil
                yeni_loglar = [l['islem'] for l in depo._loglar[log_sayisi:]]
                self.assertEqual(yeni_loglar, ['EXCEL YÜKLENDİ'])

    # --- Flask: önizleme + onay + API tutarlılığı --------------------------
    def test_arayuz_onizleme_onay_ve_api(self):
        mod, client = self.load_app()
        with self.snapshot_yamasi(), GERCEK_EXCEL.open('rb') as f:
            yanit = client.post('/yonetim/yukle', data={'_csrf_token': 'test-token', 'onizleme': '1',
                                                       'dosya': (f, GERCEK_EXCEL.name)})
        self.assertEqual(yanit.status_code, 200)
        html = yanit.get_data(as_text=True)
        self.assertIn('44 kayıt okundu', html)
        self.assertIn(f'52 satır kart sayılmadı: {MAKINE_KALINTI}', html)
        self.assertRegex(html, r'kontrol-sayi">56</span>\s*<div>\s*<strong>Kart sayılmayan satırlar')
        self.assertRegex(html, r'<span>Yeni kart</span><strong>44</strong>')
        self.assertIn('DURUM &#34;MALZEME TEDARİK&#34;: iş akışı durumu atanmaz', html)
        self.assertIn('ELDE DİZGİ: satır 67, 69, 72', html)
        self.assertIn('44 yeni kartın listesini göster', html)
        self.assertEqual(depo._kartlar, [])                     # önizleme kayıt değiştirmez
        with self.snapshot_yamasi():
            yanit = client.post('/yonetim/yukle-onay', data={'_csrf_token': 'test-token'})
        self.assertEqual(yanit.status_code, 302)
        with client.session_transaction() as oturum:
            mesaj = oturum['_flashes'][-1][1]
        self.assertIn('56 satır Talep NO/Kart Stok No içermediği için kart sayılmadı', mesaj)
        self.assertLess(len(mesaj), 1500)
        pano = client.get('/api/veriler').get_json()
        self.assertEqual(len(pano['kartlar']), 35)              # durumu boş 9 kart operasyonda görünmez
        self.assertEqual((pano['sayac']['plana_alindi'], pano['sayac']['dizgide']), (3, 5))
        teslim = client.get('/api/panel/teslimler').get_json()
        self.assertEqual(len(teslim['teslim_edilen']), 27)
        self.assertEqual({d: len(client.get(f'/api/panel/teslimler?dizgi={d}').get_json()['teslim_edilen'])
                          for d in ('MAKINE', 'ELLE', 'EUM')}, {'MAKINE': 8, 'ELLE': 13, 'EUM': 6})
        self.assertEqual(len(depo.durumu_eksik_kartlari_getir()), 9)
        # Rapor özeti, kaynakta olmayan kartları güncel sayaçlara katmaz.
        self.assertEqual(mod.ozet_hesapla()['genel']['toplam'], 44)
        self.yukle(self.kopya(lambda wb: [setattr(h, 'value', None) for h in wb['MAKİNE'][272]]))
        self.assertEqual(mod.ozet_hesapla()['genel']['plana_alindi'], 2)


    def test_arayuz_onizleme_gruplar_ve_hata_ekrani(self):
        mod, client = self.load_app()
        self.yukle()

        def onizle(path):
            with self.snapshot_yamasi(), open(path, 'rb') as f:
                return client.post('/yonetim/yukle', data={'_csrf_token': 'test-token', 'onizleme': '1',
                                                          'dosya': (f, 'plan.xlsx')})

        def degistir(wb):
            wb['MAKİNE']['W272'] = datetime(2026, 9, 10)   # NO 11 plan teslim
            wb['MAKİNE']['Y272'] = 'DİZGİDE'
            for hucre in wb['MAKİNE'][263]:               # NO 8 silindi
                hucre.value = None
        yanit = onizle(self.kopya(degistir))
        self.assertEqual(yanit.status_code, 200)
        html = yanit.get_data(as_text=True)
        self.assertIn("Excel'de artık olmayan kartlar pasifleşecek", html)
        self.assertIn('Kart güncellenecek', html)
        self.assertRegex(html, r'<tr class="onemli">\s*<td>Durum</td>\s*<td class="eski">PLANA ALINDI</td>'
                               r'\s*<td class="yeni">DİZGİDE</td>')
        self.assertRegex(html, r'<td>Plan Teslim</td>\s*<td class="eski">02.09.2026</td>\s*'
                               r'<td class="yeni">10.09.2026</td>')
        self.assertNotIn('Yeni kart oluşturulacak', html)
        self.assertNotIn('Onaylamadan önce kontrol edin.', html)   # 1/44 toplu pasifleşme değil

        # Sayfanın yarısı silinirse toplu pasifleşme uyarısı ve buton değişir.
        def yarisi(wb):
            for r in range(3, 73):
                for hucre in wb['ELDE DİZGİ'][r]:
                    hucre.value = None
        html = onizle(self.kopya(yarisi)).get_data(as_text=True)
        self.assertIn('44 aktif karttan 20 tanesi pasifleşecek', html)
        self.assertIn('Yine de uygula', html)

        # Hatalı dosya: yönetime yönlendirmek yerine sorunlar madde madde gösterilir.
        def hatali(wb):
            wb['MAKİNE']['D272'] = None
            wb['EÜM']['H9'] = None
        once = copy.deepcopy(depo._kartlar)
        yanit = onizle(self.kopya(hatali))
        self.assertEqual(yanit.status_code, 422)
        html = yanit.get_data(as_text=True)
        self.assertIn('Excel kabul edilmedi', html)
        self.assertIn('2 SORUN', html)
        self.assertIn('<span class="sorun-konum">MAKİNE · satır 272</span>', html)
        self.assertIn('<span class="sorun-konum">EÜM · satır 9</span>', html)
        self.assertEqual(depo._kartlar, once)


class ParserBirimTests(unittest.TestCase):
    """Satır sınıflandırma, hafta metni ve normalizasyon; geçici XLSX ile."""
    setUp = fixtures.SyncTests.setUp
    book = fixtures.SyncTests.book
    upload = fixtures.SyncTests.upload

    def test_hafta_metni_pazartesiye_iner(self):
        for metin, referans, beklenen in [
            ('23. Hafta (2.06 haftası)', '2025-06-11', '2025-06-02'),
            ('25. Hafta (16.06 haftası )', '2025-06-23', '2025-06-16'),
            ('34. hafta 17.08 haftası)', '2026-08-21', '2026-08-17'),
            ('37. hafta (7.09 haftası)', None, None),       # yıl: Pazartesi+hafta tutarlılığı
            ('1. hafta (29.12 haftası)', '2026-01-02', '2025-12-29'),
            ('52. hafta', '2026-12-25', '2026-12-21'),
            ('(3.06.2026 haftası)', None, '2026-06-01'),
        ]:
            with self.subTest(metin=metin):
                sonuc, uyari = ex._hafta_metni_coz(metin, referans)
                if beklenen is None:
                    self.assertEqual(date.fromisoformat(sonuc).weekday(), 0)
                    self.assertEqual(date.fromisoformat(sonuc).isocalendar()[1], 37)
                    self.assertEqual((date.fromisoformat(sonuc).day, date.fromisoformat(sonuc).month), (7, 9))
                else:
                    self.assertEqual(sonuc, beklenen)
                self.assertIsNone(uyari)
        self.assertEqual(ex._hafta_metni_coz('belirsiz', '2026-01-01'), (None, None))
        sonuc, uyari = ex._hafta_metni_coz('30. hafta (2.06 haftası)', '2025-06-11')
        self.assertEqual(sonuc, '2025-06-02')
        self.assertIn('uyuşmuyor', uyari)

    def test_eum_hafta_metni_ve_okunamayan_metin(self):
        basliklar = list(fixtures.HEADERS); basliklar[8] = 'Planlanan Başlangıç T.'
        with patch.object(fixtures, 'HEADERS', basliklar):
            sonuc = self.upload({'MAKİNE': [], 'EÜM': [
                row(1, start=None, end='2025-06-11', planned='23. Hafta (2.06 haftası)'),
                row(2, start=None, end=None, planned='belirsiz'),
                row(3, start='2026-09-16', end='2026-09-20', planned='38. hafta (14.09 haftası)'),
            ]})
        k = {x['sira']: x for x in depo._kartlar}
        self.assertEqual((k[1]['plan_baslama'], k[2]['plan_baslama'], k[3]['plan_baslama']),
                         ('2025-06-02', None, '2026-09-16'))  # açık Dizgi Başlama tarihi önceliklidir
        self.assertTrue(any('belirsiz' in u for u in sonuc['kaynak_raporu']['uyarilar']))

    def test_durum_normalizasyonu(self):
        for ham, beklenen in [('PLANA ALINDI', 'PLANA ALINDI'), ('plana alındı', 'PLANA ALINDI'),
                              (' Üretim  Plana Alındı ', 'PLANA ALINDI'), ('ÜRETİM PLANA ALINDI', 'PLANA ALINDI'),
                              ('PLANA ALİNDI', 'PLANA ALINDI'), ('Dizgide', 'DİZGİDE'),
                              ('DİZGİDE VE MALZEME BEKLENİYOR', 'DİZGİDE'), ('ÜRETİM DEVAM EDİYOR', 'DİZGİDE'),
                              ('teslim edildi', 'TESLİM EDİLDİ'), ('PDGM ÖNERİ', None), ('MALZEME TEDARİK', None)]:
            with self.subTest(ham=ham):
                self.assertEqual(ex.durum_coz(ham)[0], beklenen)
        self.assertTrue(ex.malzeme_bekliyor_mu('DİZGİDE VE MALZEME BEKLENİYOR'))

    def test_kalinti_satir_atlanir_kanitli_satir_reddedilir(self):
        kalinti = [None, None, 'STK-X', None, None, None, None, None, None, 'Sahip', 'VAR']
        numarali = [7, None, 'STK-Y', None, None, None, None, None, None, None, None]
        sonuc = self.upload({'MAKİNE': [row(1), kalinti, numarali]})
        makine = sonuc['kaynak_raporu']['sayfalar'][0]
        self.assertEqual((makine['kayit'], [a['satir'] for a in makine['atlanan']]), (1, [3, 4]))
        once = copy.deepcopy(depo._kartlar)
        for alan, deger in [(3, '5 ADET'), (7, 'DİZGİDE'), (5, '2026-09-20'), (6, '2026-09-21')]:
            with self.subTest(alan=fixtures.HEADERS[alan]):
                yarim = [None, None, 'STK-Z', None, None, None, None, None, None, None, None]
                yarim[alan] = deger
                with self.assertRaisesRegex(ex.ExcelAktarimHatasi, 'Talep NO boş, ancak satırda kayıt verisi var'):
                    self.upload({'MAKİNE': [row(1), yarim]})
                self.assertEqual(depo._kartlar, once)

    def test_excel_hata_degerleri(self):
        # Metin ve hata değerleri ' önekiyle yazılır: Excel onları yeniden yorumlamaz.
        self.assertEqual(ex._com_degerleri(((1, -2146826246), ('x', -2146826281), ('', None))),
                         ((1, "'#N/A"), ("'x", "'#DIV/0!"), ('', None)))
        self.assertEqual(ex._com_degerleri(-2146826265), (("'#REF!",),))
        self.assertEqual(ex._com_degerleri(46147.0), ((46147.0,),))
        # Kalıntı satırdaki hata değeri kayıt kanıtı değildir; kayıttaki hata reddedilir.
        sonuc = self.upload({'MAKİNE': [row(1), [None, '#N/A', 'STK', None, None, None, None, None, None, None, None]]})
        self.assertEqual(len(sonuc['kaynak_raporu']['sayfalar'][0]['atlanan']), 1)
        with self.assertRaisesRegex(ex.ExcelAktarimHatasi, 'Excel hata değeri'):
            self.upload({'MAKİNE': [row(1, end='#N/A')]})

    def test_gizli_bos_sutun_ve_gizli_satir_sentetik(self):
        path = self.book({'MAKİNE': [row(1, status='DİZGİDE'), row(2), [None] * 11, row(3)]})
        wb = openpyxl.load_workbook(path)
        ws = wb['MAKİNE']
        ws.insert_cols(4)                         # başlıksız boş sütun
        ws.column_dimensions['I'].hidden = True   # DURUM (kaymadan sonra I)
        ws.row_dimensions[3].hidden = True        # kayıt satırı gizli
        ws.row_dimensions[4].hidden = True        # boş satır gizli
        wb.save(path); wb.close()

        def snapshot(_):
            hedef = self.root / 'snapshot.xlsx'; shutil.copy2(path, hedef); return str(hedef)
        with patch.object(ex, 'excel_deger_snapshot_olustur', snapshot):
            sonuc = ex.excelden_aktar(str(path), 'tester')
        sayfa = sonuc['kaynak_raporu']['sayfalar'][0]
        self.assertEqual((sayfa['kayit'], sayfa['gizli_kayit'], sayfa['bos_satir'], sayfa['atlanan']),
                         (3, 1, 1, []))
        self.assertEqual({k['sira']: k['durum'] for k in depo._kartlar},
                         {1: 'DİZGİDE', 2: 'PLANA ALINDI', 3: 'PLANA ALINDI'})


@unittest.skipUnless(os.environ.get('PDGM_TEST_EXCEL_COM') == '1' and GERCEK_EXCEL.exists(),
                     'Set PDGM_TEST_EXCEL_COM=1 for real Excel COM')
class GercekExcelCOMTests(unittest.TestCase):
    """Aynı gerçek dosya senaryoları, snapshot'ı gerçek Microsoft Excel üretirken."""
    setUp = fixtures.SyncTests.setUp
    kopya = GercekExcelTests.kopya
    kartlar = GercekExcelTests.kartlar
    disk = GercekExcelTests.disk

    def yukle(self, path=GERCEK_EXCEL, **kw):
        return ex.excelden_aktar(str(path), 'tester', **kw)

    def test_com_orijinal_degisiklik_ve_tekrar_yukleme(self):
        beklenen = referans_kayitlar(GERCEK_EXCEL)
        sonuc = self.yukle()
        rapor = {s['kod']: s for s in sonuc['kaynak_raporu']['sayfalar']}
        self.assertEqual({k: s['kayit'] for k, s in rapor.items()}, {'MAKINE': 15, 'ELLE': 20, 'EUM': 9})
        self.assertEqual(rapor['MAKINE']['atlanan_araliklar'], MAKINE_KALINTI)
        self.assertTrue(rapor['MAKINE']['atlanan'][0]['gizli'])   # COM gizli satır bilgisini taşır
        k = self.kartlar()
        for anahtar, b in beklenen.items():
            for alan in ('talep_no', 'stok_no', 'toplam_adet', 'plan_baslama', 'plan_teslim',
                         'gerceklesen_teslim', 'durum'):
                self.assertEqual(k[anahtar][alan], b[alan], (anahtar, alan))
        once = copy.deepcopy(depo._kartlar)
        self.assertEqual(self.yukle()['degismeyen'], 44)
        self.assertEqual(depo._kartlar, once)

        def degistir(wb):
            wb['MAKİNE']['W272'] = datetime(2026, 9, 10)
            wb['MAKİNE']['Y272'] = 'DİZGİDE'
            wb['MAKİNE'].column_dimensions['Y'].hidden = True     # DURUM sütunu gizli
            for hucre in wb['MAKİNE'][263]:
                hucre.value = None                                 # NO 8 silindi
            ws = wb['EÜM']
            for s, v in zip('BCDHITUW', [10, 1900010, 'EÜM Sahip', 'EU-0000-0010', '4 ADET',
                                        '39. hafta (21.09 haftası)', datetime(2026, 9, 25), 'ÜRETİM PLANA ALINDI']):
                ws[f'{s}246'] = v
        sonuc = self.yukle(self.kopya(degistir))
        self.assertEqual((sonuc['yeni'], sonuc['guncellenen'], sonuc['pasife_alinan']), (1, 1, 1))
        k = self.kartlar()
        self.assertEqual((k[('MAKINE', 11)]['plan_teslim'], k[('MAKINE', 11)]['durum']), ('2026-09-10', 'DİZGİDE'))
        self.assertEqual(k[('MAKINE', 8)]['source_active'], 0)
        self.assertEqual((k[('EUM', 10)]['durum'], k[('EUM', 10)]['plan_baslama']), ('PLANA ALINDI', '2026-09-21'))

    def test_com_formul_hata_degeri_kayitta_reddedilir(self):
        def degistir(wb):
            wb['MAKİNE']['J272'] = '=NA()'
        with self.assertRaisesRegex(ex.ExcelAktarimHatasi, "'MAKİNE' satır 272: Excel hata değeri"):
            self.yukle(self.kopya(degistir))
        self.assertEqual(depo._kartlar, [])


if __name__ == '__main__':
    unittest.main()
