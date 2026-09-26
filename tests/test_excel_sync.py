import copy
import importlib.util
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from datetime import datetime, timezone, timedelta
from unittest.mock import patch

import openpyxl
from openpyxl.utils.datetime import to_excel, CALENDAR_MAC_1904
import depo
import excel_araclari as ex

ROOT = Path(__file__).resolve().parents[1]
HEADERS = ['NO', 'Talep NO', 'Kart Stok No', 'Kart Üretim Adet',
           'Dizgi Başlama Tarihi', 'Planlanan Teslim T.', 'Gerçekleşen Teslim T.',
           'DURUM', 'T.planlanan tarih', 'Talep Sahibi', 'PCB']


def row(no, qty=5, status='PLANA ALINDI', start='2026-09-14', end='2026-09-20',
        actual=None, planned=None, stock='ABC', request='123'):
    return [no, request, stock, qty, start, end, actual, status, planned, 'Owner', 'PCB']


class SyncTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        values = {'VERI_KLASORU': str(self.root), 'YEDEK_KLASORU': str(self.root/'backups'),
                  'KARTLAR_DOSYA': str(self.root/'kartlar.xlsx'),
                  'LOG_DOSYA': str(self.root/'log.xlsx'),
                  'YUKLEME_DOSYA': str(self.root/'uploads.xlsx'),
                  '_kartlar': [], '_loglar': [], '_yuklemeler': []}
        for name, value in values.items():
            p = patch.object(depo, name, value); p.start(); self.addCleanup(p.stop)
        depo.kur()

    def book(self, sheets, epoch=None):
        wb = openpyxl.Workbook(); wb.remove(wb.active)
        if epoch: wb.epoch = epoch
        for name, rows in sheets.items():
            ws = wb.create_sheet(name); ws.append(HEADERS)
            for r in rows: ws.append(r)
        path = self.root/'input.xlsx'; wb.save(path); wb.close()
        return path

    def upload(self, sheets, epoch=None):
        path = self.book(sheets, epoch)
        # COM is the external boundary; actual parser, merge and XLSX persistence run.
        def snapshot(_):
            target = self.root/'snapshot.xlsx'; shutil.copy2(path, target); return str(target)
        with patch.object(ex, 'excel_deger_snapshot_olustur', snapshot):
            return ex.excelden_aktar(str(path), 'tester')

    def cards(self):
        return {k['sira']: k for k in depo.kartlari_getir(False)}

    def test_same_upload_idempotent(self):
        sheets = {'MAKİNE': [row(1), row(2)]}
        self.upload(sheets); before = self.cards()
        result = self.upload(sheets)
        self.assertEqual(result['yeni'], 0)
        self.assertEqual([(k['id'], k['durum']) for k in self.cards().values()],
                         [(k['id'], k['durum']) for k in before.values()])
        self.assertEqual(len(depo._kartlar), 2)

    def test_duplicate_date_change_keeps_identity(self):
        a,b = row(1),row(2,end='2026-09-21')
        self.upload({'MAKİNE':[a,b]}); old = self.cards()[2]['id']
        b[5] = '2026-09-24'; self.upload({'MAKİNE':[a,b]})
        self.assertEqual(len(depo._kartlar), 2)
        self.assertEqual(self.cards()[2]['id'], old)
        self.assertEqual(self.cards()[2]['plan_teslim'], '2026-09-24')
        self.assertEqual(self.cards()[1]['plan_teslim'], '2026-09-20')

    def test_start_date_change_keeps_identity(self):
        a,b = row(1,end=None),row(2,end=None,start='2026-09-15')
        self.upload({'MAKİNE':[a,b]}); old = self.cards()[2]['id']
        b[4]='2026-09-16'; self.upload({'MAKİNE':[a,b]})
        self.assertEqual(len(depo._kartlar),2); self.assertEqual(self.cards()[2]['id'],old)
        self.assertEqual(self.cards()[2]['plan_baslama'],'2026-09-16')

    def test_status_sequence_including_delivered_without_date(self):
        r=row(1)
        for status in ['PLANA ALINDI','DİZGİDE','TESLİM EDİLDİ']:
            r[7]=status; self.upload({'MAKİNE':[r]})
            self.assertEqual(depo.kartlari_getir()[0]['durum'],status)
        self.assertIsNone(self.cards()[1]['gerceklesen_teslim'])
        self.assertEqual(self.cards()[1]['tamamlanan_adet'],5)

    def test_actual_date_clear_and_change(self):
        r=row(1,status='TESLİM EDİLDİ',actual='2026-09-18')
        self.upload({'MAKİNE':[r]})
        for value in ['2026-09-19',None]:
            r[6]=value; self.upload({'MAKİNE':[r]})
            self.assertEqual(self.cards()[1]['gerceklesen_teslim'],value)

    def test_delete_only_one_duplicate_and_disallow_operation(self):
        a,b=row(1,status='DİZGİDE'),row(2)
        self.upload({'MAKİNE':[a,b]}); old=self.cards()[1]['id']
        self.upload({'MAKİNE':[b]})
        self.assertEqual([k['sira'] for k in depo.kartlari_getir()],[2])
        self.assertEqual(depo.kart_getir(old)['source_active'],0)
        with self.assertRaises(depo.IsKuralHatasi):
            depo.kart_bitir(old,1,'tester','admin','makine')

    def test_three_rows_reorder_insert_delete(self):
        a,b,c=row(1),row(2,qty=2),row(3,qty=8)
        self.upload({'MAKİNE':[a,b,c]}); ids={n:k['id'] for n,k in self.cards().items()}
        self.upload({'MAKİNE':[c,row(4,qty=9),a,b]})
        self.assertEqual({n:self.cards()[n]['id'] for n in ids},ids)
        self.upload({'MAKİNE':[c,a]})
        self.assertEqual({k['sira'] for k in depo.kartlari_getir()},{1,3})
        self.assertEqual(self.cards()[3]['toplam_adet'],8)

    def test_machine_hand_and_eum_same_no_independent(self):
        sheets={'MAKİNE':[row(1,status='DİZGİDE')],
                'ELDE DİZGİ':[row(1,qty=2,status='TESLİM EDİLDİ',actual='2026-09-18')],
                'EÜM':[row(1,qty=9)]}
        self.upload(sheets)
        before={k['dizgi_kod']:k for k in depo.kartlari_getir()}
        self.assertEqual((before['ELLE']['toplam_adet'],before['ELLE']['durum']),(2,'TESLİM EDİLDİ'))
        self.assertEqual((before['MAKINE']['toplam_adet'],before['MAKINE']['durum']),(5,'DİZGİDE'))
        sheets['MAKİNE']=[]; self.upload(sheets)
        after={k['dizgi_kod']:k for k in depo.kartlari_getir()}
        self.assertNotIn('MAKINE',after)
        self.assertEqual(after['ELLE']['id'],before['ELLE']['id'])
        self.assertEqual(after['EUM']['id'],before['EUM']['id'])

    def test_restart_preserves_source_identity_and_deletion(self):
        self.upload({'MAKİNE':[row(1),row(2)]})
        self.upload({'MAKİNE':[row(2,end='2026-09-25')]})
        before=copy.deepcopy(depo._kartlar)
        depo._kartlar=[]; depo.kur()
        self.assertEqual([(k['id'],k.get('source_key'),k['source_active']) for k in depo._kartlar],
                         [(k['id'],k.get('source_key'),k['source_active']) for k in before])
        self.assertEqual([k['sira'] for k in depo.kartlari_getir()],[2])

    def test_blank_or_invalid_status_preserves_manual_progress_notes(self):
        r=row(1,status='DİZGİDE'); self.upload({'MAKİNE':[r]})
        k=depo._kartlar[0]; k.update(tamamlanan_adet=2,operator='worker',aciklama='history')
        # Boş DURUM ve bilerek durumsuz bırakılan kaynak durumu iş akışını korur.
        # Tanınmayan diğer yazımlar artık importu durdurur (test_onceki_hatalar).
        for status in [None,'MALZEME TEDARİK']:
            r[7]=status; r[5]='2026-09-25'; result=self.upload({'MAKİNE':[r]})
            k=self.cards()[1]
            self.assertEqual((k['durum'],k['tamamlanan_adet'],k['operator'],k['aciklama']),
                             ('DİZGİDE',2,'worker','history'))
            self.assertEqual(result['workflow_korundu'],1)

    def test_eum_planned_date_monday(self):
        for value,want in [('17.09.2026','2026-09-14'),('20.09.2026','2026-09-14'),
                           ('21.09.2026','2026-09-21'),('01.01.2026','2025-12-29'),
                           (datetime(2026,9,17,23,59),'2026-09-14'),
                           (to_excel(datetime(2026,9,17)),'2026-09-14')]:
            with self.subTest(value=value):
                self.upload({'MAKİNE':[], 'EÜM':[row(1,start=None,end=None,planned=value)]})
                self.assertEqual(self.cards()[1]['plan_baslama'],want)

    def test_eum_existing_planned_start_header(self):
        headers=list(HEADERS); headers[8]='Planlanan Başlangıç T.'
        with patch.object(sys.modules[__name__],'HEADERS',headers):
            self.upload({'MAKİNE':[], 'EÜM':[row(1,start=None,planned='17.09.2026')]})
        self.assertEqual(self.cards()[1]['plan_baslama'],'2026-09-14')

    def test_1904_numeric_dates(self):
        serial=to_excel(datetime(2026,9,17),epoch=CALENDAR_MAC_1904)
        self.upload({'MAKİNE':[], 'EÜM':[row(1,start=None,end=None,planned=serial)]},CALENDAR_MAC_1904)
        self.assertEqual(self.cards()[1]['plan_baslama'],'2026-09-14')

    def test_missing_or_duplicate_no_rejects_without_changes(self):
        self.upload({'MAKİNE':[row(1)]}); before=copy.deepcopy(depo._kartlar)
        for rows in [[row(None)],[row(1),row(1,stock='OTHER')]]:
            with self.assertRaises((ex.ExcelAktarimHatasi,depo.VeriDogrulamaHatasi)):
                self.upload({'MAKİNE':rows})
            self.assertEqual(depo._kartlar,before)

    def test_valid_empty_source_deactivates_all(self):
        self.upload({'MAKİNE':[row(1)]}); self.upload({'MAKİNE':[]})
        self.assertEqual(depo.kartlari_getir(),[])
        self.assertEqual(len(depo._kartlar),1)

    def test_manual_card_not_hijacked(self):
        manual=depo.admin_kart_ekle('123','ABC',10,'tester')
        self.upload({'MAKİNE':[row(1)]})
        self.assertEqual(depo.kart_getir(manual['id'])['kaynak'],'MANUEL')
        self.assertEqual(depo.kart_getir(manual['id'])['toplam_adet'],10)
        self.assertEqual(len(depo._kartlar),2)

    def test_replace_failure_rolls_back_memory_and_disk(self):
        self.upload({'MAKİNE':[row(1)]})
        before=copy.deepcopy(depo._kartlar)
        paths=[depo.KARTLAR_DOSYA,depo.LOG_DOSYA,depo.YUKLEME_DOSYA]
        disk={p:Path(p).read_bytes() for p in paths}; original=os.replace
        def failing(src,dst):
            if str(src).endswith('.yeni') and dst==depo.LOG_DOSYA:
                raise OSError('injected second replace failure')
            return original(src,dst)
        with patch.object(depo.os,'replace',failing):
            with self.assertRaises(OSError): self.upload({'MAKİNE':[row(1,end='2026-09-25')]})
        self.assertEqual(depo._kartlar,before)
        self.assertEqual({p:Path(p).read_bytes() for p in paths},disk)

    def legacy(self, rows, no=True):
        self.upload({'MAKİNE':rows})
        for index,k in enumerate(depo._kartlar):
            for field in ['source_key','source_row_id','source_sheet','legacy_anahtar']:
                k.pop(field,None)
            k['anahtar']='123|ABC'+(f'#{index+1}' if index else '')
            if not no: k['sira']=None
            k['aciklama']=f'original note {index+1}'
            k['operator']='original operator'
        # Write a real old-schema workbook without the newly added columns.
        schema=[(h,a) for h,a in depo.KART_ALANLARI if not a.startswith('source_') or a=='source_active']
        schema=[(h,a) for h,a in schema if a!='legacy_anahtar']
        depo._yaz(depo.KARTLAR_DOSYA,schema,depo._kartlar,'Kartlar')
        depo.kur()
        return copy.deepcopy(depo._kartlar)

    def test_legacy_migration_preserves_ids_notes_and_backup(self):
        old=self.legacy([row(1),row(2,qty=2)])
        result=self.upload({'MAKİNE':[row(2,qty=2,end='2026-09-25'),row(1)]})
        self.assertEqual(result['yeni'],0)
        self.assertEqual(self.cards()[2]['id'],old[1]['id'])
        self.assertEqual(self.cards()[2]['aciklama'],'original note 2')
        self.assertEqual(self.cards()[2]['operator'],'original operator')
        self.assertEqual(self.cards()[2]['legacy_anahtar'],'123|ABC#2')
        self.assertTrue((Path(result['yedek'])/'kartlar.xlsx').exists())
        depo.kur()
        self.assertEqual(self.cards()[2]['source_row_id'],'NO:2')
        self.assertEqual(self.upload({'MAKİNE':[row(1),row(2)]})['yeni'],0)

    def test_legacy_unique_without_no_migrates(self):
        old=self.legacy([row(1)],no=False)
        self.upload({'MAKİNE':[row(10,end='2026-09-25')]})
        self.assertEqual(self.cards()[10]['id'],old[0]['id'])
        self.assertEqual(self.cards()[10]['aciklama'],'original note 1')

    def test_ambiguous_legacy_duplicate_rejected_atomically(self):
        before=self.legacy([row(1),row(2)],no=False)
        disk=Path(depo.KARTLAR_DOSYA).read_bytes()
        with self.assertRaisesRegex(depo.VeriDogrulamaHatasi,'belirsiz'):
            self.upload({'MAKİNE':[row(2),row(1)]})
        self.assertEqual(depo._kartlar,before)
        self.assertEqual(Path(depo.KARTLAR_DOSYA).read_bytes(),disk)

    def test_explicit_row_id_without_no_survives_reorder(self):
        headers=list(HEADERS);headers[0]='PDGM_ROW_ID'
        with patch.object(sys.modules[__name__],'HEADERS',headers):
            self.upload({'MAKİNE':[row('a'),row('b',qty=2)]})
            ids={k['source_row_id']:k['id'] for k in depo._kartlar}
            self.upload({'MAKİNE':[row('b',qty=2,end='2026-09-25'),row('a')]})
        self.assertEqual({k['source_row_id']:k['id'] for k in depo._kartlar},ids)

    def test_timezone_dates_keep_source_calendar_day(self):
        for raw in ['2026-09-21T00:30:00+03:00','2026-09-21T23:30:00-05:00']:
            self.upload({'MAKİNE':[],'EÜM':[row(1,start=None,end=None,planned=raw)]})
            self.assertEqual(self.cards()[1]['plan_baslama'],'2026-09-21')
        self.assertEqual(depo.tarih_coz(datetime(2026,9,21,0,30,tzinfo=timezone(timedelta(hours=3)))),'2026-09-21')

    def test_invalid_date_aborts_all_rows(self):
        self.upload({'MAKİNE':[row(1)]}); before=copy.deepcopy(depo._kartlar)
        for value in ['31.02.2026','not a date',True]:
            with self.assertRaises(ex.ExcelAktarimHatasi):
                self.upload({'MAKİNE':[row(1,end='2026-09-25')],
                             'EÜM':[row(2,start=None,end=None,planned=value)]})
            self.assertEqual(depo._kartlar,before)

    def test_plan_fields_clear_without_creating_new_card(self):
        self.upload({'MAKİNE':[row(1)]})
        self.upload({'MAKİNE':[row(1,start=None,end=None)]})
        self.assertEqual(len(depo._kartlar),1)
        self.assertIsNone(self.cards()[1]['plan_baslama'])
        self.assertIsNone(self.cards()[1]['plan_teslim'])

    def test_valid_status_updates_quantity_but_keeps_notes_and_hidden(self):
        self.upload({'MAKİNE':[row(1,status='DİZGİDE')]})
        depo._kartlar[0].update(tamamlanan_adet=2,aciklama='history',admin_gizli=1)
        self.upload({'MAKİNE':[row(1,qty=7,status='TESLİM EDİLDİ')]})
        k=self.cards()[1]
        self.assertEqual((k['tamamlanan_adet'],k['aciklama'],k['admin_gizli']),(7,'history',1))
        self.upload({'MAKİNE':[row(1,qty=3,status='PLANA ALINDI')]})
        self.assertEqual(self.cards()[1]['tamamlanan_adet'],0)

    def test_manual_partial_quantity_conflict_rolls_back(self):
        self.upload({'MAKİNE':[row(1,status='DİZGİDE')]})
        depo._kartlar[0]['tamamlanan_adet']=3
        before=copy.deepcopy(depo._kartlar)
        with self.assertRaises(depo.IsKuralHatasi):
            self.upload({'MAKİNE':[row(1,qty=2,status=None)]})
        self.assertEqual(depo._kartlar,before)

    def test_blank_status_cleared_actual_date_not_replaced_by_old_timestamp(self):
        self.upload({'MAKİNE':[row(1,status='TESLİM EDİLDİ',actual='2026-09-18')]})
        depo._kartlar[0]['teslim_zamani']='2026-09-18 12:00:00'
        self.upload({'MAKİNE':[row(1,status=None)]})
        self.assertIsNone(self.cards()[1]['gerceklesen_teslim'])
        self.assertIsNone(self.cards()[1]['sapma'])

    def test_backups_do_not_overwrite_within_same_second(self):
        with patch.object(depo,'simdi',return_value='2026-09-19 12:00:00'):
            a=self.upload({'MAKİNE':[row(1)]})['yedek']
            b=self.upload({'MAKİNE':[row(1,end='2026-09-25')]})['yedek']
        self.assertNotEqual(a,b)

    def test_direct_import_rejects_duplicate_source_key(self):
        self.upload({'MAKİNE':[row(1)]}); before=copy.deepcopy(depo._kartlar)
        wb=openpyxl.load_workbook(self.root/'input.xlsx')
        parsed=[];ex._sayfa_satirlarini_coz(wb.active,'MAKİNE',depo.DIZGI_TIPI_MAKINE,{},parsed);wb.close()
        with self.assertRaises(depo.VeriDogrulamaHatasi):
            depo.excel_import_uygula('test.xlsx','tester',parsed+parsed)
        self.assertEqual(depo._kartlar,before)

    def test_malformed_populated_row_is_not_a_deletion(self):
        self.upload({'MAKİNE':[row(1)]}); before=copy.deepcopy(depo._kartlar)
        for col in [1,2]:
            bad=row(1);bad[col]=None
            with self.assertRaises(ex.ExcelAktarimHatasi): self.upload({'MAKİNE':[bad]})
            self.assertEqual(depo._kartlar,before)

    def test_lossy_legacy_no_rejected_instead_of_losing_history(self):
        before=self.legacy([row('001')])
        for rows in [[row('001')],[row('1'),row('001')]]:
            with self.assertRaises(depo.VeriDogrulamaHatasi): self.upload({'MAKİNE':rows})
            self.assertEqual(depo._kartlar,before)

    def test_fresh_python_process_reads_same_persisted_cards(self):
        self.upload({'MAKİNE':[row(1),row(2)]})
        self.upload({'MAKİNE':[row(2,end='2026-09-25',status='TESLİM EDİLDİ')]})
        code='''import sys,os,json,depo
root=sys.argv[1]
depo.VERI_KLASORU=root
depo.YEDEK_KLASORU=os.path.join(root,'backups')
depo.KARTLAR_DOSYA=os.path.join(root,'kartlar.xlsx')
depo.LOG_DOSYA=os.path.join(root,'log.xlsx')
depo.YUKLEME_DOSYA=os.path.join(root,'uploads.xlsx')
depo.kur()
print(json.dumps([(k['id'],k['source_row_id'],k['plan_teslim'],k['durum']) for k in depo.kartlari_getir()]))
'''
        result=subprocess.run([sys.executable,'-B','-c',code,str(self.root)],cwd=ROOT,
                              capture_output=True,text=True,check=True)
        self.assertEqual(json.loads(result.stdout),[[2,'NO:2','2026-09-25','TESLİM EDİLDİ']])

    def test_concurrent_commits_keep_each_snapshot_whole(self):
        from concurrent.futures import ThreadPoolExecutor
        from threading import Barrier
        parsed=[]
        for suffix in ['A','B']:
            path=self.book({'MAKİNE':[row(1,stock=suffix+'1'),row(2,stock=suffix+'2')]})
            wb=openpyxl.load_workbook(path); rows=[]
            ex._sayfa_satirlarini_coz(wb.active,'MAKİNE',depo.DIZGI_TIPI_MAKINE,{},rows)
            wb.close();parsed.append(rows)
        barrier=Barrier(2)
        def commit(rows):
            barrier.wait();return depo.excel_import_uygula('thread.xlsx','tester',rows)
        with ThreadPoolExecutor(max_workers=2) as pool: list(pool.map(commit,parsed))
        stocks={k['stok_no'] for k in depo.kartlari_getir()}
        self.assertIn(stocks,[{'A1','A2'},{'B1','B2'}])
        depo.kur()
        self.assertEqual({k['stok_no'] for k in depo.kartlari_getir()},stocks)

    def test_admin_note_edit_does_not_invent_imported_event_dates(self):
        for status,completed in [('TESLİM EDİLDİ',5),('DİZGİDE',0)]:
            self.upload({'MAKİNE':[row(1,status=status)]})
            card=self.cards()[1]
            depo.admin_kart_duzenle(card['id'],status,completed,5,'updated note','tester',
                                   gerceklesen_teslim='')
            updated=self.cards()[1]
            self.assertEqual(updated['aciklama'],'updated note')
            for field in ['gerceklesen_teslim','baslama_zamani','bitis_zamani','teslim_zamani']:
                self.assertIsNone(updated[field],field)

    def test_inactive_notes_are_admin_only(self):
        self.upload({'MAKİNE':[row(1)]}); card=self.cards()[1]
        self.upload({'MAKİNE':[]})
        with self.assertRaises(depo.IsKuralHatasi):
            depo.kart_not_guncelle(card['id'],'stale note','worker','operator')
        depo.kart_not_guncelle(card['id'],'audit note','admin','admin')
        self.assertIn('audit note',depo.kart_getir(card['id'])['aciklama'])

    def test_manual_workflow_still_records_real_events(self):
        self.upload({'MAKİNE':[row(1)]}); card=self.cards()[1]
        depo.kart_baslat(card['id'],5,'worker','operator','makine')
        depo.kart_bitir(card['id'],2,'worker','operator','makine')
        with self.assertRaises(depo.IsKuralHatasi):
            depo.kart_teslim_et(card['id'],'worker','operator','makine')
        depo.kart_bitir(card['id'],3,'worker','operator','makine')
        depo.kart_teslim_et(card['id'],'worker','operator','makine')
        delivered=self.cards()[1]
        self.assertEqual(delivered['durum'],'TESLİM EDİLDİ')
        for field in ['gerceklesen_teslim','baslama_zamani','bitis_zamani','teslim_zamani']:
            self.assertIsNotNone(delivered[field])


if __name__=='__main__': unittest.main()
