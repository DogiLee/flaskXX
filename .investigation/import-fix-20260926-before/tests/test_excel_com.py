"""Opt-in integration against installed Microsoft Excel; only temporary files."""
import os
import unittest
import openpyxl
from openpyxl.utils.datetime import CALENDAR_MAC_1904, to_excel
from datetime import datetime
import depo
import excel_araclari as ex
from tests import test_excel_sync as fixtures


@unittest.skipUnless(os.environ.get('PDGM_TEST_EXCEL_COM')=='1', 'Set PDGM_TEST_EXCEL_COM=1 for real Excel COM')
class ExcelCOMTests(unittest.TestCase):
    setUp=fixtures.SyncTests.setUp
    book=fixtures.SyncTests.book

    def test_text_identifiers_keep_leading_zeros_and_long_digits(self):
        path=self.book({'MAKİNE':[fixtures.row('001'),fixtures.row('1'),
            fixtures.row('12345678901234567890'),fixtures.row('N001')]})
        wb=openpyxl.load_workbook(path)
        wb['MAKİNE'].column_dimensions['A'].hidden=True
        for cell in list(wb['MAKİNE']['A'])[1:]: cell.number_format='@'
        wb.save(path);wb.close()
        ex.excelden_aktar(str(path),'tester')
        before={k['source_row_id']:k['id'] for k in depo._kartlar}
        self.assertEqual(set(before),{'NO:001','NO:1','NO:12345678901234567890','NO:N001'})
        ex.excelden_aktar(str(path),'tester')
        self.assertEqual(before,{k['source_row_id']:k['id'] for k in depo._kartlar})

    def test_real_snapshot_hidden_identity_1904_and_reupload(self):
        path=self.book({'MAKİNE':[fixtures.row(1)],'EÜM':[
            fixtures.row(1,start=None,end=None,planned=to_excel(datetime(2026,9,17),CALENDAR_MAC_1904))
        ]},epoch=CALENDAR_MAC_1904)
        wb=openpyxl.load_workbook(path)
        for ws in wb:
            ws.column_dimensions['A'].hidden=True
            ws.column_dimensions['K'].hidden=True
        wb.save(path);wb.close()
        ex.excelden_aktar(str(path),'tester')
        before={k['dizgi_kod']:k for k in depo.kartlari_getir()}
        self.assertEqual(before['MAKINE']['source_row_id'],'NO:1')
        self.assertEqual(before['EUM']['plan_baslama'],'2026-09-14')
        self.assertFalse(before['MAKINE'].get('pcb'))
        wb=openpyxl.load_workbook(path)
        wb['MAKİNE']['F2']='2026-09-25'
        wb['MAKİNE']['H2']='TESLİM EDİLDİ'
        wb.save(path);wb.close()
        ex.excelden_aktar(str(path),'tester')
        after={k['dizgi_kod']:k for k in depo.kartlari_getir()}
        self.assertEqual(after['MAKINE']['id'],before['MAKINE']['id'])
        self.assertEqual(after['MAKINE']['durum'],'TESLİM EDİLDİ')
        self.assertEqual(after['MAKINE']['plan_teslim'],'2026-09-25')


if __name__=='__main__': unittest.main()
