"""Gerçek PDGM örnek Excel'inden kapsamlı elle test paketi üretir.

Kullanım (proje kökünden):
    .venv\\Scripts\\python.exe tools\\test_paketi_uret.py [hedef_klasor]

Kaynak dosyanın yapısı (sayfalar, başlıklar, gizli sütun/satırlar, kalıntı
satırlar, hücre stilleri) korunur; yeni satırlar mevcut kayıt satırlarının
stiliyle eklenir. Tarihler üretildiği güne göre hesaplanır; rozet beklentileri
(BUGÜN BAŞLAMALI, SON GÜN, SÜRE AŞILDI...) o gün için geçerlidir. Başka bir gün
test edilecekse paketi o gün yeniden üretin.

Dosyalar ve sıra için bkz. TEST_REHBERI.md (bu betik aynı klasöre yazar).
"""
from __future__ import annotations

import copy
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

KOK = Path(__file__).resolve().parents[1]
KAYNAK = KOK / "PDGM_Kart_dizgi_Talepleri_Üretim_Takvimi.xlsx"

SAHIPLER = ["Gökhan Uz", "Sarper Van", "Doğu Çan", "Ali Tek", "Ayşe Al", "Volkan Batı", "Tümer Ak",
            "Ayla Bil", "Emre Kibar", "Yiğit Alp", "Mustafa As", "Orkun Kan", "Burak Tan", "Deniz Var"]
SORUMLULAR = ["Ahmet Düzgün", "Mehmet Büyük"]


def _gun(bugun, fark):
    return datetime.combine(bugun + timedelta(days=fark), datetime.min.time()) if fark is not None else None


def _hafta_metni(gun: date) -> str:
    pazartesi = gun - timedelta(days=gun.weekday())
    return f"{pazartesi.isocalendar()[1]}. hafta ({pazartesi.day}.{pazartesi.month:02d} haftası)"


# --- Kayıt tanımları ---------------------------------------------------------
# bas/tes/ger: bugüne göre gün farkı (None = boş). adet metin veya sayı.
# "not": test rehberinde bu satırın neyi sınadığı.
MAKINE = [
    dict(no=16, talep=1900016, adet="4 ADET", bas=7, tes=10, durum="PLANA ALINDI", not_="Plana alındı, ileri tarihli"),
    dict(no=17, talep=1900017, adet="3 ADET", bas=0, tes=3, durum="PLANA ALINDI", not_="BUGÜN BAŞLAMALI rozeti"),
    dict(no=18, talep=1900018, adet="6 ADET", bas=-5, tes=2, durum="PLANA ALINDI", not_="BAŞLAMADI (+5 gün) rozeti"),
    dict(no=19, talep=1900019, adet="2 ADET", bas=14, tes=18, durum="plana alındı", not_="küçük harf DURUM normalizasyonu"),
    dict(no=20, talep=1900020, adet="10 ADET", bas=-3, tes=5, durum="DİZGİDE", not_="PLANINDA; operatör kısmi adet girecek"),
    dict(no=21, talep=1900021, adet="2 ADET", bas=-2, tes=0, durum="DİZGİDE", not_="SON GÜN rozeti"),
    dict(no=22, talep=1900022, adet="3 ADET", bas=-4, tes=1, durum="DİZGİDE", not_="SON 1 GÜN rozeti"),
    dict(no=23, talep=1900023, adet="5 ADET", bas=-10, tes=-3, durum="DİZGİDE", not_="SÜRE AŞILDI (3 gün)"),
    dict(no=24, talep=1900024, adet="12 ADET", bas=-1, tes=6, durum="DİZGİDE", not_="operatör bitirip teslim edecek"),
    dict(no=25, talep=1900025, adet="4 ADET", bas=-20, tes=-15, ger=-15, durum="TESLİM EDİLDİ", not_="ZAMANINDA TESLİM"),
    dict(no=26, talep=1900026, adet="4 ADET", bas=-20, tes=-15, ger=-12, durum="TESLİM EDİLDİ", not_="GEÇ TESLİM (+3 gün)"),
    dict(no=27, talep=1900027, adet="1 ADET", bas=-10, tes=-8, ger=-8, durum="Teslim Edildi ", not_="yazım farkı + boşluk"),
    dict(no=28, talep=1900028, adet="2 ADET", bas=-9, tes=-6, durum="TESLİM EDİLDİ", not_="teslim tarihi boş uyarısı"),
    dict(no=29, talep=1900029, adet="8 ADET", bas=None, tes=None, durum=None, not_="DURUM boş: durumu eksik listesi"),
    dict(no=30, talep=1900030, adet="3 ADET", bas=5, tes=9, ger=-1, durum="PLANA ALINDI", not_="tarih dolu ama teslim değil uyarısı"),
    dict(no=31, talep=1900031, stok="AD-T031-0001", adet="3 ADET", bas=-15, tes=-10, ger=-10, durum="TESLİM EDİLDİ", not_="aynı Talep+Stok 1. parti"),
    dict(no=32, talep=1900031, stok="AD-T031-0001", adet="4 ADET", bas=-2, tes=4, durum="DİZGİDE", not_="aynı Talep+Stok 2. parti"),
    dict(no=33, talep=1900031, stok="AD-T031-0001", adet="5 ADET", bas=10, tes=14, durum="PLANA ALINDI", not_="aynı Talep+Stok 3. parti"),
    dict(no=34, talep=1900034, stok="AD-T034-0001", adet="5 ADET", bas=-2, tes=4, durum="DİZGİDE", not_="7 adet: 5 makine (bu) + 2 elle (ELDE NO 28)"),
    dict(no=35, talep=1900035, adet="1.500 ADET", bas=20, tes=30, durum="PLANA ALINDI", not_="binlik ayraçlı adet"),
    dict(no=36, talep=1900036, adet=12, bas=21, tes=25, durum="PLANA ALINDI", not_="sayısal adet"),
    dict(no=37, talep=1900037, adet="6 ADET", bas=2, tes=6, durum="PLANA ALINDI", gizli=True, not_="gizli satırdaki kayıt"),
    dict(no=38, talep=1900038, adet="2 ADET", bas=3, tes="metin", durum="PLANA ALINDI", not_="metin tarih (gg.aa.yyyy)"),
    dict(no=39, talep=1900039, adet="2 ADET", bas=-1, tes=3, durum="DİZGİDE", not_="operatör teslim edecek; Excel geride kalacak"),
    dict(no=40, talep=1900040, adet="10 ADET", bas=0, tes=9, durum="PLANA ALINDI", not_="operatör dizgiye alacak; Excel geride kalacak"),
    dict(no=41, talep=1900041, adet="2 ADET", bas=-8, tes=-4, ger=-4, durum="TESLİM EDİLDİ", not_="2. adımda Excel'de DİZGİDE'ye geri alınacak"),
    dict(no=42, talep=1900042, adet="3 ADET", bas=-1, tes=5, durum="DİZGİDE", not_="2. adımda NO başka talebe verilecek"),
    dict(no=43, talep=1900043, adet="2 ADET", bas=8, tes=12, durum="PLANA ALINDI", not_="2. adımda satır silinecek"),
    dict(no=44, talep=1900044, adet="2 ADET", bas=9, tes=13, durum="PLANA ALINDI", not_="2. adımda kalıntı bırakılarak silinecek"),
    dict(no=45, talep=1900045, adet="7 ADET", bas=-3, tes=8, durum="DİZGİDE", pcb="VAR", not_="PCB bilgisi"),
]
ELDE = [
    dict(no=21, talep=1910021, adet="2 ADET", bas=2, tes=5, durum="DİZGİ İÇİN BEKLİYOR", not_="PLANA'ya iner; 2. adımda DİZGİDE"),
    dict(no=22, talep=1910022, adet="3 ADET", bas=0, tes=2, durum="DİZGİ İÇİN BEKLİYOR", not_="elle operatörü dizgiye alacak"),
    dict(no=23, talep=1910023, adet="4 ADET", bas=-3, tes=3, durum="DİZGİDE VE MALZEME BEKLENİYOR", not_="DİZGİDE + malzeme bekliyor"),
    dict(no=24, talep=1910024, adet="2 ADET", bas=-2, tes=4, durum="DİZGİDE VE MALZEME BEKLİYOR", not_="diğer yazım"),
    dict(no=25, talep=1910025, adet="5 ADET", bas=None, tes=None, durum="MALZEME TEDARİK", not_="durumsuz + malzeme; 2. adımda PLANA"),
    dict(no=26, talep=1910026, adet="6 ADET", bas=-2, tes=3, durum="DİZGİDE", not_="elle dizgide"),
    dict(no=27, talep=1910027, adet="2 ADET", bas=-14, tes=-10, ger=-10, durum="TESLİM EDİLDİ", not_="elle teslim"),
    dict(no=28, talep=1900034, stok="AD-T034-0001", adet="2 ADET", bas=-6, tes=-3, ger=-3, durum="TESLİM EDİLDİ", not_="7 adetin 2 elle kısmı (bitmiş)"),
    dict(no=29, talep=1910029, adet="4 ADET", bas=1, tes=4, durum="DİZGİ İÇİN BEKLİYOR", not_="PLANA"),
    dict(no=30, talep=1910030, adet="6 ADET", bas=-1, tes=5, durum="DİZGİDE", not_="elle operatörü kısmi adet girecek"),
    dict(no=31, talep=1910031, adet="3 ADET", bas=-12, tes=-9, ger=-6, durum="TESLİM EDİLDİ", not_="GEÇ TESLİM"),
    dict(no=32, talep=1910032, adet="2 ADET", bas=-9, tes=-2, durum="DİZGİDE", not_="SÜRE AŞILDI"),
    dict(no=33, talep=1910033, adet="1 ADET", bas=4, tes=7, durum="DİZGİ İÇİN BEKLİYOR", not_="2. adımda silinecek"),
    dict(no=34, talep=1910034, adet="2 ADET", bas=None, tes=None, durum=None, not_="DURUM boş"),
    dict(no=35, talep=1910035, adet="3 ADET", bas=-1, tes=4, durum="dizgide", not_="küçük harf"),
]
EUM = [
    dict(no=10, talep=1920010, adet="4 ADET", hafta=7, tes=11, durum="PLANA ALINDI", not_="yeni EÜM değeri; 2. adımda DİZGİDE"),
    dict(no=11, talep=1920011, adet="3 ADET", hafta=0, tes=4, durum="ÜRETİM PLANA ALINDI", not_="EÜM operatörü dizgiye alacak"),
    dict(no=12, talep=1920012, adet="2 ADET", hafta=14, tes=18, durum="üretim  plana alındı ", not_="küçük harf + fazla boşluk"),
    dict(no=13, talep=1920013, adet="5 ADET", hafta=-3, tes=4, durum="ÜRETİM DEVAM EDİYOR", not_="DİZGİDE"),
    dict(no=14, talep=1920014, adet="2 ADET", hafta=10, tes=14, durum="EMTD PLANLANDI", not_="PLANA"),
    dict(no=15, talep=1920015, adet="3 ADET", hafta=-12, tes=-8, ger=-7, durum="TESLİM EDİLDİ", not_="EÜM teslim"),
    dict(no=16, talep=1920016, adet="4 ADET", hafta=21, tes=None, durum="PDGM ÖNERİ", not_="durumsuz; 2. adımda PLANA"),
    dict(no=17, talep=1920017, adet="2 ADET", hafta=28, tes=33, durum="PLANA ALINDI", hafta_bicim="yalniz_hafta", not_="yalnız 'NN. hafta'"),
    dict(no=18, talep=1920018, adet="2 ADET", hafta="belirsiz", tes=40, durum="PLANA ALINDI", not_="okunamayan hafta metni uyarısı"),
    dict(no=19, talep=1920019, adet="1 ADET", hafta=35, tes=40, durum="PLANA ALINDI", hafta_bicim="tarih", not_="hafta sütununda tarih"),
    dict(no=20, talep=1920020, adet="3 ADET", hafta=-9, tes=-2, durum="ÜRETİM DEVAM EDİYOR", not_="SÜRE AŞILDI"),
    dict(no=21, talep=1900016, stok="AD-T016-EUM", adet="2 ADET", hafta=7, tes=12, durum="PLANA ALINDI", not_="MAKİNE NO 16 ile aynı Talep, ayrı kart"),
    dict(no=22, talep=1920022, adet="6 ADET", hafta=-2, tes=5, durum="ÜRETİM DEVAM EDİYOR", not_="EÜM dizgide"),
    dict(no=23, talep=1920023, adet="2 ADET", hafta=3, tes=8, durum="ÜRETİM PLANA ALINDI", not_="EÜM plan"),
]
YENI_MAKINE = [  # 2. adımda eklenecek
    dict(no=46, talep=1900046, adet="2 ADET", bas=6, tes=9, durum="PLANA ALINDI", not_="yeni kayıt"),
    dict(no=47, talep=1900047, adet="3 ADET", bas=-1, tes=4, durum="DİZGİDE", not_="yeni kayıt"),
    dict(no=48, talep=1900048, adet="1 ADET", bas=-6, tes=-2, ger=-2, durum="TESLİM EDİLDİ", not_="yeni kayıt"),
]


# --- Yazma yardımcıları ---------------------------------------------------------
def _stil_kopyala(ws, kaynak_satir, hedef_satir, son_sutun):
    for sutun in range(1, son_sutun + 1):
        ws.cell(hedef_satir, sutun)._style = copy.copy(ws.cell(kaynak_satir, sutun)._style)


def _makine_yaz(ws, r, k, bugun, sablon=272):
    _stil_kopyala(ws, sablon, r, 26)
    bas = _gun(bugun, k.get("bas"))
    tes = (_gun(bugun, k["tes"]) if k.get("tes") != "metin"
           else (bugun + timedelta(days=10)).strftime("%d.%m.%Y"))
    stok = k.get("stok") or f"AD-T{k['no']:03d}-0001"
    sira = k["no"]
    degerler = {
        "C": k["no"], "D": k["talep"], "E": SAHIPLER[sira % len(SAHIPLER)], "F": _gun(bugun, -30),
        "G": SORUMLULAR[sira % 2], "H": "01", "I": stok, "J": k["adet"],
        "K": f"MP-T{sira:03d}", "L": "01", "M": "HBT", "N": _gun(bugun, -60),
        # Gizli sipariş bloğundaki aynı adlı "Planlanan Teslim T." kasıtlı olarak farklı:
        # yanlış sütun okunursa tarihler hemen fark edilir.
        "O": _gun(bugun, -45), "P": _gun(bugun, -44), "Q": "MAKİNE", "R": _gun(bugun, -25), "S": _gun(bugun, -25),
        "T": _hafta_metni(bas.date()) if bas else None, "U": bas, "W": tes,
        "X": _gun(bugun, k.get("ger")), "Y": k.get("durum"), "Z": k.get("pcb"),
    }
    for harf, deger in degerler.items():
        ws[f"{harf}{r}"] = deger
    ws.row_dimensions[r].hidden = bool(k.get("gizli"))


def _elde_yaz(ws, r, k, bugun, sablon=70):
    _stil_kopyala(ws, sablon, r, 23)
    stok = k.get("stok") or f"EL-T{k['no']:03d}-0001"
    degerler = {
        "B": k["no"], "C": k["talep"], "D": SAHIPLER[(k["no"] + 3) % len(SAHIPLER)], "E": _gun(bugun, -30),
        "F": SORUMLULAR[k["no"] % 2], "G": "01", "H": stok, "I": k["adet"],
        "L": "HBT", "M": _gun(bugun, -60), "N": _gun(bugun, -45), "O": _gun(bugun, -44), "P": "MANUEL",
        "Q": _gun(bugun, -25), "R": _gun(bugun, -20),
        "S": _gun(bugun, k.get("bas")), "T": _gun(bugun, k.get("tes")), "U": _gun(bugun, k.get("ger")),
        "V": k.get("durum"), "W": k.get("malzeme"),
    }
    for harf, deger in degerler.items():
        ws[f"{harf}{r}"] = deger
    ws.row_dimensions[r].hidden = False


def _eum_yaz(ws, r, k, bugun, sablon=244):
    _stil_kopyala(ws, sablon, r, 24)
    stok = k.get("stok") or f"EU-T{k['no']:03d}-0001"
    if k["hafta"] == "belirsiz":
        hafta = "belirsiz"
    else:
        gun = bugun + timedelta(days=k["hafta"])
        bicim = k.get("hafta_bicim")
        if bicim == "yalniz_hafta":
            hafta = f"{gun.isocalendar()[1]}. hafta"
        elif bicim == "tarih":
            hafta = gun.strftime("%d.%m.%Y")
        else:
            hafta = _hafta_metni(gun)
    degerler = {
        "B": k["no"], "C": k["talep"], "D": SAHIPLER[(k["no"] + 7) % len(SAHIPLER)], "E": _gun(bugun, -30),
        "F": SORUMLULAR[k["no"] % 2], "G": "01", "H": stok, "I": k["adet"], "L": "YURTDIŞI",
        "M": _gun(bugun, -60), "N": _gun(bugun, -45), "O": _gun(bugun, -44), "P": "ÜRETİMDE",
        "Q": _gun(bugun, -20), "R": _gun(bugun, -18), "S": _gun(bugun, -15),
        "T": hafta, "U": _gun(bugun, k.get("tes")), "V": _gun(bugun, k.get("ger")),
        "W": k.get("durum"), "X": k.get("malzeme"),
    }
    for harf, deger in degerler.items():
        ws[f"{harf}{r}"] = deger
    ws.row_dimensions[r].hidden = False


def _satir_bul(ws, no_sutunu, no, baslangic):
    for r in range(baslangic, ws.max_row + 1):
        if ws[f"{no_sutunu}{r}"].value == no and ws[f"{'D' if no_sutunu == 'C' else 'C'}{r}"].value:
            return r
    raise KeyError(f"{ws.title} NO {no}")


# --- Dosya setleri ---------------------------------------------------------------
def ilk_yukleme(bugun):
    wb = openpyxl.load_workbook(KAYNAK)
    ws = wb["MAKİNE"]
    r = 279
    for i, k in enumerate(MAKINE):
        _makine_yaz(ws, r, k, bugun)
        r += 1
        if i == 9:          # tamamen boş satır
            r += 1
        if i == 19:         # yalnız boşluk içeren satır ve gizli boş satır
            for harf in "DEIJY":
                ws[f"{harf}{r}"] = "   "
            ws.row_dimensions[r + 1].hidden = True
            r += 2
    ws[f"I{r + 1}"] = "AD-KALINTI-0001"        # görünür kalıntı: yalnız stok
    ws[f"C{r + 2}"], ws[f"G{r + 2}"] = 99, "Ahmet Düzgün"   # kalıntı: NO + sorumlu

    ws = wb["ELDE DİZGİ"]
    for i, k in enumerate(ELDE):
        _elde_yaz(ws, 73 + i, k, bugun)

    ws = wb["EÜM"]
    for i, k in enumerate(EUM):
        _eum_yaz(ws, 246 + i, k, bugun)
    return wb


def guncelleme(bugun):
    """2. adım: tarih/durum/adet değişiklikleri, yeni/silinen satırlar, NO'nun başka talebe
    verilmesi, Excel'de geri alma, gizli DURUM sütunu, görünür yinelenen başlık."""
    wb = ilk_yukleme(bugun)
    ws = wb["MAKİNE"]
    bul = lambda no: _satir_bul(ws, "C", no, 205)
    ws[f"W{bul(16)}"] = _gun(bugun, 12)                                   # tarih değişikliği
    ws[f"Y{bul(18)}"] = "DİZGİDE"                                         # Excel ilerletti
    ws[f"W{bul(20)}"] = _gun(bugun, 7)                                    # DİZGİDE kalır, tarih değişti
    ws[f"Y{bul(23)}"], ws[f"X{bul(23)}"] = "TESLİM EDİLDİ", _gun(bugun, -1)
    ws[f"J{bul(36)}"] = 15                                                 # adet değişikliği
    ws[f"Y{bul(37)}"] = "DİZGİDE"                                         # gizli satırda güncelleme
    ws[f"Y{bul(41)}"], ws[f"X{bul(41)}"] = "DİZGİDE", None                 # Excel'de geri alma
    r42 = bul(42)                                                         # NO 42 yeni talebe verildi
    ws[f"D{r42}"], ws[f"I{r42}"], ws[f"J{r42}"], ws[f"Y{r42}"] = 1909042, "AD-T042-YENI", "5 ADET", "PLANA ALINDI"
    ws[f"U{r42}"], ws[f"W{r42}"] = _gun(bugun, 3), _gun(bugun, 8)
    for hucre in ws[bul(43)]:                                            # satır tamamen silindi
        hucre.value = None
    r44 = bul(44)                                                         # kalıntı bırakılarak silindi
    for harf in "DEFHJKLMNOPQRSTUWXY":
        ws[f"{harf}{r44}"] = None
    son = max(r for r in range(279, ws.max_row + 1) if ws[f"C{r}"].value)
    for i, k in enumerate(YENI_MAKINE):
        _makine_yaz(ws, son + 3 + i, k, bugun)
    ws.column_dimensions["Y"].hidden = True       # DURUM sütunu gizli: yine okunmalı
    ws.column_dimensions["O"].hidden = False      # aynı adlı ikinci "Planlanan Teslim T." görünür

    ws = wb["ELDE DİZGİ"]
    bul = lambda no: _satir_bul(ws, "B", no, 3)
    ws[f"V{bul(21)}"] = "DİZGİDE"
    ws[f"V{bul(25)}"], ws[f"S{bul(25)}"], ws[f"T{bul(25)}"] = "DİZGİ İÇİN BEKLİYOR", _gun(bugun, 3), _gun(bugun, 6)
    for hucre in ws[bul(33)]:
        hucre.value = None

    ws = wb["EÜM"]
    bul = lambda no: _satir_bul(ws, "B", no, 3)
    ws[f"W{bul(10)}"] = "ÜRETİM DEVAM EDİYOR"
    ws[f"W{bul(16)}"], ws[f"U{bul(16)}"] = "PLANA ALINDI", _gun(bugun, 26)
    _eum_yaz(ws, 246 + len(EUM), dict(no=26, talep=1920026, adet="3 ADET", hafta=5, tes=10,
                                      durum="ÜRETİM PLANA ALINDI"), bugun)
    return wb


def excel_guncel(bugun):
    """4. adım: Excel operatörün uygulamada yaptıklarına yetişti; karar gerekmez."""
    wb = guncelleme(bugun)
    ws = wb["MAKİNE"]
    bul = lambda no: _satir_bul(ws, "C", no, 205)
    ws[f"Y{bul(24)}"], ws[f"X{bul(24)}"] = "TESLİM EDİLDİ", _gun(bugun, 0)
    ws[f"Y{bul(39)}"], ws[f"X{bul(39)}"] = "TESLİM EDİLDİ", _gun(bugun, 0)
    ws[f"Y{bul(40)}"] = "DİZGİDE"
    wb["ELDE DİZGİ"][f"V{_satir_bul(wb['ELDE DİZGİ'], 'B', 22, 3)}"] = "DİZGİDE"
    wb["EÜM"][f"W{_satir_bul(wb['EÜM'], 'B', 11, 3)}"] = "ÜRETİM DEVAM EDİYOR"
    return wb


def hatali_dosyalar(bugun):
    """Reddedilmesi gereken (H) ve uyarı gösterip vazgeçilecek (U) dosyalar."""
    def temel():
        return excel_guncel(bugun)

    def makine_satiri(wb, no):
        return _satir_bul(wb["MAKİNE"], "C", no, 205)

    dosyalar = {}
    wb = temel(); wb["MAKİNE"][f"D{makine_satiri(wb, 16)}"] = None
    dosyalar["H01_EKSIK_TALEP_NO"] = wb
    wb = temel(); wb["MAKİNE"][f"Y{makine_satiri(wb, 16)}"] = "TESLİM EDİLDİ (KISMİ)"
    dosyalar["H02_TANINMAYAN_DURUM"] = wb
    wb = temel(); wb["MAKİNE"][f"C{makine_satiri(wb, 17)}"] = 16
    dosyalar["H03_TEKRARLANAN_NO"] = wb
    wb = temel(); wb["MAKİNE"][f"J{makine_satiri(wb, 16)}"] = "on adet"
    dosyalar["H04_GECERSIZ_ADET"] = wb
    wb = temel(); wb["MAKİNE"][f"W{makine_satiri(wb, 16)}"] = "31.02.2026"
    dosyalar["H05_GECERSIZ_TARIH"] = wb
    wb = temel(); r = makine_satiri(wb, 16); wb["MAKİNE"][f"U{r}"] = _gun(bugun, 30)
    dosyalar["H06_BASLAMA_TESLIMDEN_SONRA"] = wb
    wb = temel()
    wb["MAKİNE"][f"D{makine_satiri(wb, 16)}"] = None
    wb["ELDE DİZGİ"][f"I{_satir_bul(wb['ELDE DİZGİ'], 'B', 21, 3)}"] = "on adet"
    wb["EÜM"][f"W{_satir_bul(wb['EÜM'], 'B', 10, 3)}"] = "Bitti"
    dosyalar["H07_UC_HATA_BIRDEN"] = wb
    wb = temel(); wb["MAKİNE"].title = "URETIM"
    dosyalar["H08_MAKINE_SAYFASI_YOK"] = wb
    wb = temel(); wb["MAKİNE"]["V4"] = "NO"; wb["MAKİNE"].column_dimensions["V"].hidden = False
    dosyalar["H09_IKI_NO_SUTUNU"] = wb
    wb = temel(); del wb["EÜM"]
    dosyalar["U01_EUM_SAYFASI_YOK"] = wb
    wb = temel(); ws = wb["ELDE DİZGİ"]
    for r in range(3, ws.max_row + 1):
        for hucre in ws[r]:
            hucre.value = None
    dosyalar["U02_ELDE_SAYFASI_BOSALTILDI"] = wb
    return dosyalar


def uret(hedef: Path, bugun: date | None = None) -> dict[str, Path]:
    bugun = bugun or date.today()
    hedef.mkdir(parents=True, exist_ok=True)
    sirali = {
        "01_ILK_YUKLEME": ilk_yukleme(bugun),
        "02_GUNCELLEME": guncelleme(bugun),
        "03_AYNI_DOSYA_TEKRAR": guncelleme(bugun),
        "04_EXCEL_YETISTI": excel_guncel(bugun),
        **hatali_dosyalar(bugun),
    }
    yollar = {}
    for ad, wb in sirali.items():
        yol = hedef / f"{ad}.xlsx"
        wb.save(yol)
        wb.close()
        yollar[ad] = yol
    return yollar


if __name__ == "__main__":
    hedef = Path(sys.argv[1]) if len(sys.argv) > 1 else KOK / "outputs" / f"pdgm-test-paketi-{date.today():%Y%m%d}"
    for ad, yol in uret(hedef).items():
        print(f"{ad:32} {yol}")
