#!/usr/bin/env python3
"""macOS / COM yokken test Excel'ini ayni parser + depo.import ile uygular.

Uretim yolu (Windows COM snapshot) degil; sadece ortak test icin.
Kullanim (proje kokunden):

  .venv/bin/python test_verileri/mac_import_yardimci.py test_verileri/PDGM_TEST_INPUT.xlsx
"""

from __future__ import annotations

import os
import sys

KOK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, KOK)

import openpyxl

import depo
import excel_araclari as ex


def aktar_com_siz(dosya_yolu: str, kullanici: str = "admin"):
    wb = openpyxl.load_workbook(dosya_yolu, data_only=True)
    parsed = []
    anahtar_gruplari = {}
    uyari_sayisi = 0
    sayfa_bulundu = False

    try:
        for sayfa_adi, dizgi_tipi, zorunlu in ex.KAYNAK_SAYFALARI:
            if sayfa_adi not in wb.sheetnames:
                if zorunlu:
                    raise ex.ExcelAktarimHatasi(
                        f"'{sayfa_adi}' sayfası bulunamadı. "
                        f"Dosyadaki sayfalar: {', '.join(wb.sheetnames)}"
                    )
                continue
            sayfa_bulundu = True
            uyari_sayisi += ex._sayfa_satirlarini_coz(
                wb[sayfa_adi], sayfa_adi, dizgi_tipi, anahtar_gruplari, parsed
            )
        if not sayfa_bulundu:
            raise ex.ExcelAktarimHatasi("Excel'de işlenecek kart satırı bulunamadı.")
    finally:
        wb.close()

    return depo.excel_import_uygula(
        dosya_adi=os.path.basename(dosya_yolu),
        kullanici=kullanici,
        satirlar=parsed,
        uyari_sayisi=uyari_sayisi,
    )


def main():
    if len(sys.argv) < 2:
        print("Kullanim: python test_verileri/mac_import_yardimci.py <xlsx>")
        sys.exit(2)

    yol = os.path.abspath(sys.argv[1])
    if not os.path.isfile(yol):
        print(f"Dosya yok: {yol}")
        sys.exit(1)

    depo.kur()
    sonuc = aktar_com_siz(yol)
    print("Import OK")
    for k, v in sonuc.items():
        if k == "pasife_listesi":
            print(f"  {k}: {len(v)} ornek")
            continue
        print(f"  {k}: {v}")


if __name__ == "__main__":
    main()
